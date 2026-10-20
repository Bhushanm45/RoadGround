import os
import pptx
from pptx.util import Inches, Pt

def enlarge_presentation():
    input_path = 'C:/Users/aryan/.gemini/antigravity/brain/0c03362b-d615-4f91-8096-b1c9b36d96de/.user_uploaded/media_1792226287834_06596d45.pptx'
    output_workspace_path = 'C:/Users/aryan/.gemini/antigravity/scratch/RoadGuard-AI/RoadGuard_AI_Fonts_Enlarged.pptx'
    output_downloads_path = 'C:/Users/aryan/Downloads/RoadGuard_AI_Fonts_Enlarged.pptx'
    
    prs = pptx.Presentation(input_path)
    print(f"Loaded presentation with {len(prs.slides)} slides.")
    
    # SLIDE 1: MUST REMAIN COMPLETELY UNCHANGED
    print("Slide 1 left completely unchanged.")
    
    # =========================================================================
    # SLIDE 2: Idea Title
    # =========================================================================
    s2 = prs.slides[1]
    for sh in s2.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                if "Idea Title:" in p.text:
                    for r in p.runs:
                        if r.font and r.font.size:
                            r.font.size = Pt(18.5)
                        elif hasattr(p, 'font') and p.font and p.font.size:
                            r.font.size = Pt(18.5)
                    print("  Slide 2: Enlarged Idea Title banner text to 18.5pt.")

    # =========================================================================
    # SLIDE 3: Technical Approach
    # =========================================================================
    s3 = prs.slides[2]
    for sh in s3.shapes:
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text.strip()
        
        # 1. Hardware Column Header
        if "EXPECTED PROTOTYPE SETUP" in txt:
            for p in sh.text_frame.paragraphs:
                p.font.size = Pt(13.0)
            print("  Slide 3: Enlarged Prototype Setup Header to 13.0pt.")
            
        # 2. Hardware Specs Box below photo
        elif "Sensor:" in txt or "Microcontroller:" in txt:
            for p in sh.text_frame.paragraphs:
                p.space_after = Pt(2.0)
                for r in p.runs:
                    if r.font and r.font.bold:
                        r.font.size = Pt(10.5)
                    else:
                        r.font.size = Pt(10.0)
            print("  Slide 3: Enlarged Hardware Specs bullets to 10.5pt bold / 10.0pt regular.")
            
        # 3. Waveforms Header
        elif "VIBRATION SIGNATURE ANALYSIS" in txt:
            for p in sh.text_frame.paragraphs:
                p.font.size = Pt(12.5)
            print("  Slide 3: Enlarged Waveforms Header to 12.5pt.")
            
        # 4. Waveform Discrimination Header
        elif "ILLUSTRATIVE WAVEFORM DISCRIMINATION" in txt or "WAVEFORM DISCRIMINATION" in txt:
            for p in sh.text_frame.paragraphs:
                p.font.size = Pt(11.5)
            print("  Slide 3: Enlarged Waveform Criteria Header to 11.5pt.")
            
        # 5. Waveform Criteria Bullets
        elif "Pothole (Cavity):" in txt or "Speed Breaker (Hump):" in txt:
            for p in sh.text_frame.paragraphs:
                p.space_after = Pt(2.0)
                for r in p.runs:
                    if "Calibration Note" in p.text:
                        if r.font and r.font.bold:
                            r.font.size = Pt(9.5)
                        else:
                            r.font.size = Pt(9.0)
                    else:
                        if r.font and r.font.bold:
                            r.font.size = Pt(10.5)
                        else:
                            r.font.size = Pt(10.0)
            print("  Slide 3: Enlarged Waveform Criteria bullets to 10.5pt bold / 10.0pt regular.")
            
        # 6. ML Workflow Header
        elif "5-STEP ML WORKFLOW" in txt:
            for p in sh.text_frame.paragraphs:
                p.font.size = Pt(12.5)
            print("  Slide 3: Enlarged ML Workflow Header to 12.5pt.")
            
        # 7. Prototype Configuration Note
        elif "Prototype Configuration Note:" in txt:
            for idx, p in enumerate(sh.text_frame.paragraphs):
                if idx == 0:
                    for r in p.runs:
                        r.font.size = Pt(11.0)
                    if not p.runs:
                        p.font.size = Pt(11.0)
                else:
                    for r in p.runs:
                        r.font.size = Pt(10.0)
                    if not p.runs:
                        p.font.size = Pt(10.0)
            print("  Slide 3: Enlarged Prototype Configuration Note to 11.0pt / 10.0pt.")

    # =========================================================================
    # SLIDE 4: Feasibility and Viability
    # =========================================================================
    s4 = prs.slides[3]
    for sh in s4.shapes:
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text.strip()
        
        # Column Headers
        if "1. FEASIBILITY" in txt or "2. POTENTIAL CHALLENGES" in txt or "3. MITIGATION STRATEGIES" in txt:
            for p in sh.text_frame.paragraphs:
                p.font.size = Pt(13.5)
            print(f"  Slide 4: Enlarged Column Header '{txt[:20]}' to 13.5pt.")
            
        # Three Columns Body Text
        elif "[ ✔ ]" in txt or "[ ! ]" in txt or "[ ★ ]" in txt or "[ ? ]" in txt:
            for p in sh.text_frame.paragraphs:
                p.space_after = Pt(10.0)
                for r in p.runs:
                    if r.font and r.font.bold:
                        r.font.size = Pt(12.8)
                    else:
                        r.font.size = Pt(11.8)
            print("  Slide 4: Enlarged Column Bullets to 12.8pt bold / 11.8pt regular.")
            
        # Bottom Summary Banner
        elif "Feasibility Summary:" in txt:
            for p in sh.text_frame.paragraphs:
                p.font.size = Pt(14.0)
                for r in p.runs:
                    r.font.size = Pt(14.0)
            print("  Slide 4: Enlarged Bottom Feasibility Summary Banner to 14.0pt.")

    # =========================================================================
    # SLIDE 5: Impact and Benefits
    # =========================================================================
    s5 = prs.slides[4]
    for sh in s5.shapes:
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text.strip()
        
        # Impact Category Headers
        if "SOCIAL IMPACT" in txt or "ECONOMIC IMPACT" in txt or "ENVIRONMENTAL & OPERATIONAL IMPACT" in txt:
            for p in sh.text_frame.paragraphs:
                p.font.size = Pt(12.0)
            print(f"  Slide 5: Enlarged Impact Header '{txt[:20]}' to 12.0pt.")
            
        # Impact Bullets
        elif "Safer Urban Commutes:" in txt or "Reduces Survey Expenditure:" in txt or "Zero Dedicated Survey Emissions:" in txt:
            for p in sh.text_frame.paragraphs:
                p.space_after = Pt(3.0)
                for r in p.runs:
                    if r.font and r.font.bold:
                        r.font.size = Pt(12.2)
                    elif r.font and r.font.size and r.font.size.pt > 10.5:
                        r.font.size = Pt(12.2)
                    else:
                        r.font.size = Pt(11.4)
            print("  Slide 5: Enlarged Impact bullets to 12.2pt bold / 11.4pt regular.")
            
        # Future Scope Roadmap Header
        elif "FUTURE SCOPE & SYSTEM ROADMAP" in txt:
            for p in sh.text_frame.paragraphs:
                p.font.size = Pt(11.8)
            print("  Slide 5: Enlarged Future Scope Header to 11.8pt.")
            
        # Future Scope Bullets below flowchart
        elif "Proposed Stages:" in txt or "Objective: Enable" in txt:
            for p in sh.text_frame.paragraphs:
                p.space_after = Pt(3.0)
                for r in p.runs:
                    if "Proposed Stages:" in p.text:
                        r.font.size = Pt(10.8)
                    else:
                        r.font.size = Pt(10.2)
            print("  Slide 5: Enlarged Future Scope roadmap description to 10.8pt / 10.2pt.")
            
        # Dashboard Demo Badge
        elif "ILLUSTRATIVE DEMO CONCEPT DATA" in txt:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10.5)
            print("  Slide 5: Enlarged Dashboard Demo Badge to 10.5pt.")
            
        # Dashboard Header
        elif "PROPOSED MUNICIPAL GIS CONCEPT" in txt:
            for p in sh.text_frame.paragraphs:
                p.font.size = Pt(11.5)
            print("  Slide 5: Enlarged Dashboard Header to 11.5pt.")
            
        # Dashboard Disclaimer Body
        elif "Smart-City Fleet Monitoring Concept" in txt or "Illustrative Demo Data Notice:" in txt:
            for idx, p in enumerate(sh.text_frame.paragraphs):
                if idx == 0:
                    for r in p.runs:
                        r.font.size = Pt(11.5)
                    if not p.runs:
                        p.font.size = Pt(11.5)
                else:
                    for r in p.runs:
                        r.font.size = Pt(10.2)
                    if not p.runs:
                        p.font.size = Pt(10.2)
            print("  Slide 5: Enlarged Dashboard Disclaimer body to 11.5pt / 10.2pt.")

    # =========================================================================
    # SLIDE 6: Research and References
    # =========================================================================
    s6 = prs.slides[5]
    for sh in s6.shapes:
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text.strip()
        
        # Reference Card Headers
        if "1. HARDWARE SPECIFICATION" in txt or "2. MACHINE LEARNING FRAMEWORK" in txt or \
           "3. VALIDATION METHODOLOGY" in txt or "4. PEER-REVIEWED RESEARCH" in txt:
            for p in sh.text_frame.paragraphs:
                p.font.size = Pt(13.5)
            print(f"  Slide 6: Enlarged Reference Header '{txt[:25]}' to 13.5pt.")
            
        # Cards 1, 2, 3 Body (Primary Source, Technical Scope, Project Role)
        elif "Primary Source:" in txt or "Technical Scope:" in txt or "Project Role:" in txt:
            for p in sh.text_frame.paragraphs:
                p.space_after = Pt(7.0)
                for r in p.runs:
                    if r.font and r.font.bold:
                        r.font.size = Pt(14.0)
                    else:
                        r.font.size = Pt(13.2)
            print("  Slide 6: Enlarged Reference Cards 1-3 to 14.0pt bold / 13.2pt regular.")
            
        # Card 4 Body (Peer-Reviewed Papers)
        elif "Mednis et al." in txt or "De Zoysa et al." in txt or "Mohan et al." in txt:
            for p in sh.text_frame.paragraphs:
                p.space_after = Pt(7.0)
                for r in p.runs:
                    r.font.size = Pt(13.2)
            print("  Slide 6: Enlarged Reference Card 4 literature to 13.2pt.")
            
        # Bottom Summary Card
        elif "PROJECT SUMMARY & CORE INNOVATION" in txt or "RoadGuard AI combines low-cost vibration sensing" in txt:
            for idx, p in enumerate(sh.text_frame.paragraphs):
                if "PROJECT SUMMARY" in p.text:
                    for r in p.runs:
                        r.font.size = Pt(14.5)
                    if not p.runs:
                        p.font.size = Pt(14.5)
                elif "RoadGuard AI combines" in p.text:
                    for r in p.runs:
                        r.font.size = Pt(18.5)
                    if not p.runs:
                        p.font.size = Pt(18.5)
                elif "Hardware Hackathon Track" in p.text:
                    for r in p.runs:
                        r.font.size = Pt(13.0)
                    if not p.runs:
                        p.font.size = Pt(13.0)
            print("  Slide 6: Enlarged Bottom Project Summary quote to 18.5pt / 14.5pt / 13.0pt.")

    # Save to workspace and downloads
    prs.save(output_workspace_path)
    print(f"\nSuccessfully saved enlarged presentation to: {output_workspace_path}")
    
    prs.save(output_downloads_path)
    print(f"Successfully copied enlarged presentation to: {output_downloads_path}")

if __name__ == '__main__':
    enlarge_presentation()
