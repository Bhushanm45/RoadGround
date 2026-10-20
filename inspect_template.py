import pptx
from pptx import Presentation
import os

prs = Presentation('C:/Users/aryan/.gemini/antigravity/scratch/RoadGuard-AI/template.pptx')
print(f"Total slides: {len(prs.slides)}")
print(f"Slide width: {prs.slide_width.inches} inches, height: {prs.slide_height.inches} inches")

for idx, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {idx + 1} ====================")
    for s_idx, shape in enumerate(slide.shapes):
        shape_type = str(shape.shape_type)
        has_text = shape.has_text_frame
        text = ""
        if has_text:
            text = "\n".join([p.text for p in shape.text_frame.paragraphs if p.text.strip()])
        elif shape.has_table:
            text = "[TABLE]"
            rows_txt = []
            for r in shape.table.rows:
                rows_txt.append(" | ".join([c.text.strip() for c in r.cells]))
            text += "\n" + "\n".join(rows_txt)
        
        print(f"Shape {s_idx}: '{shape.name}', type={shape_type}, left={shape.left.inches:.2f}in, top={shape.top.inches:.2f}in, w={shape.width.inches:.2f}in, h={shape.height.inches:.2f}in")
        if text:
            print(f"   Content:\n   " + text.replace("\n", "\n   "))
