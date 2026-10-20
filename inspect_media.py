import pptx
from pptx import Presentation

prs = Presentation('C:/Users/aryan/.gemini/antigravity/scratch/RoadGuard-AI/template.pptx')

print("Slide master count:", len(prs.slide_masters))
for sm_idx, sm in enumerate(prs.slide_masters):
    print(f"Master {sm_idx}: shapes={len(sm.shapes)}, layouts={len(sm.slide_layouts)}")
    for s in sm.shapes:
        print(f"  Master shape: {s.name}, {s.shape_type}")

# Check images embedded in pptx
import zipfile
with zipfile.ZipFile('C:/Users/aryan/.gemini/antigravity/scratch/RoadGuard-AI/template.pptx', 'r') as z:
    media = [f for f in z.namelist() if f.startswith('ppt/media/')]
    print("\nMedia files in template:")
    for m in media:
        info = z.getinfo(m)
        print(f"  {m} ({info.file_size} bytes)")
