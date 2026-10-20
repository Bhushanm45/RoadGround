import pptx

uploaded_path = 'C:/Users/aryan/.gemini/antigravity/brain/0c03362b-d615-4f91-8096-b1c9b36d96de/.user_uploaded/media_1792226287834_06596d45.pptx'
prs = pptx.Presentation(uploaded_path)

for s_idx in range(1, len(prs.slides)):
    s = prs.slides[s_idx]
    print(f'==================== SLIDE {s_idx+1} ====================')
    for sh_idx, sh in enumerate(s.shapes):
        if sh.has_text_frame:
            tf = sh.text_frame
            print(f'\n--- Shape {sh_idx}: {sh.name} (type={sh.shape_type}, left={sh.left/914400:.2f}", top={sh.top/914400:.2f}", w={sh.width/914400:.2f}", h={sh.height/914400:.2f}") ---')
            l_m = tf.margin_left.inches if tf.margin_left else 0
            r_m = tf.margin_right.inches if tf.margin_right else 0
            t_m = tf.margin_top.inches if tf.margin_top else 0
            b_m = tf.margin_bottom.inches if tf.margin_bottom else 0
            print(f'    Margins: L={l_m:.2f}", R={r_m:.2f}", T={t_m:.2f}", B={b_m:.2f}", wrap={tf.word_wrap}')
            for p_idx, p in enumerate(tf.paragraphs):
                p_txt = p.text.strip().encode('ascii', errors='replace').decode('ascii')
                if not p_txt:
                    continue
                runs_info = []
                for r in p.runs:
                    sz = f'{r.font.size.pt:.1f}pt' if r.font and r.font.size else 'None'
                    bld = 'B' if r.font and r.font.bold else 'R'
                    name = r.font.name if r.font else 'None'
                    clean_r = r.text.encode('ascii', errors='replace').decode('ascii')
                    runs_info.append(f'"{clean_r[:20]}"({sz},{bld},{name})')
                sb = p.space_before.pt if p.space_before else 0
                sa = p.space_after.pt if p.space_after else 0
                print(f'    P{p_idx} (sb={sb:.1f}, sa={sa:.1f}): {p_txt[:60]}...')
                print(f'       Runs: {runs_info}')
        elif sh.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            print(f'\n--- Picture {sh_idx}: {sh.name} (left={sh.left/914400:.2f}", top={sh.top/914400:.2f}", w={sh.width/914400:.2f}", h={sh.height/914400:.2f}") ---')
