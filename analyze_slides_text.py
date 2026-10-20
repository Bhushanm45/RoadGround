import pptx

uploaded_path = 'C:/Users/aryan/.gemini/antigravity/brain/0c03362b-d615-4f91-8096-b1c9b36d96de/.user_uploaded/media_1792226287834_06596d45.pptx'
prs = pptx.Presentation(uploaded_path)

for s_idx in range(1, len(prs.slides)):
    s = prs.slides[s_idx]
    print(f'==================== SLIDE {s_idx+1} ====================')
    for sh_idx, sh in enumerate(s.shapes):
        if sh.has_text_frame:
            tf = sh.text_frame
            txt_lines = []
            for p in tf.paragraphs:
                t = p.text.strip().encode('ascii', errors='replace').decode('ascii')
                if t:
                    sizes = [f"{r.font.size.pt:.1f}" if (r.font and r.font.size) else (f"def:{p.font.size.pt:.1f}" if (hasattr(p, 'font') and p.font and p.font.size) else "none") for r in p.runs]
                    txt_lines.append(f"  P: '{t[:70]}' (sizes={sizes})")
            if txt_lines:
                print(f"Shape {sh_idx}: {sh.name} (left={sh.left/914400:.2f}\", top={sh.top/914400:.2f}\", w={sh.width/914400:.2f}\", h={sh.height/914400:.2f}\")")
                for l in txt_lines:
                    print(l)
