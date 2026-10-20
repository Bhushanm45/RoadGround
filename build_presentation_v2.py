import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def remove_shape(shape):
    sp = shape._element
    sp.getparent().remove(sp)

def add_header_card(slide, left, top, width, header_height, total_height, 
                    title_text, header_bg, body_bg, border_color, title_size=11.5):
    # Main container
    container = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(total_height))
    container.fill.solid()
    container.fill.fore_color.rgb = body_bg
    container.line.color.rgb = border_color
    container.line.width = Pt(1.5)
    
    # Header bar
    header = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(header_height))
    header.fill.solid()
    header.fill.fore_color.rgb = header_bg
    header.line.color.rgb = header_bg
    header.line.width = Pt(1.0)
    
    tf = header.text_frame
    tf.word_wrap = False
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = "Arial"
    p.font.size = Pt(title_size)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    # Body text frame
    body_box = slide.shapes.add_textbox(Inches(left + 0.08), Inches(top + header_height + 0.04), 
                                        Inches(width - 0.16), Inches(total_height - header_height - 0.08))
    tf_body = body_box.text_frame
    tf_body.word_wrap = True
    tf_body.margin_left = Inches(0.12)
    tf_body.margin_right = Inches(0.12)
    tf_body.margin_top = Inches(0.06)
    tf_body.margin_bottom = Inches(0.06)
    return tf_body

def main():
    template_path = "C:/Users/aryan/.gemini/antigravity/scratch/RoadGuard-AI/template.pptx"
    output_pptx = "C:/Users/aryan/.gemini/antigravity/scratch/RoadGuard-AI/RoadGuard_AI_Hardware_Hackathon_Idea.pptx"
    
    prs = Presentation(template_path)
    print(f"Loaded template with {len(prs.slides)} slides.")
    
    # Assets
    asset_dir = "C:/Users/aryan/.gemini/antigravity/scratch/RoadGuard-AI/assets"
    pmpml_bus_img = os.path.join(asset_dir, "pmpml_bus.jpg")
    sensor_img = os.path.join(asset_dir, "mpu6050_sensor.jpg")
    proto_placeholder_img = os.path.join(asset_dir, "prototype_placeholder.png")
    dashboard_img = os.path.join(asset_dir, "dark_blue_dashboard.png")
    comparison_cards_img = os.path.join(asset_dir, "comparison_cards.png")
    flowchart_img = os.path.join(asset_dir, "flowchart.png")
    waveforms_img = os.path.join(asset_dir, "waveforms.png")
    s2_concept_flowchart_img = os.path.join(asset_dir, "slide2_concept_flowchart.png")
    s3_ml_flowchart_img = os.path.join(asset_dir, "slide3_ml_flowchart.png")
    s5_roadmap_flowchart_img = os.path.join(asset_dir, "slide5_roadmap_flowchart.png")
    
    # Colors
    c_navy_dark = RGBColor(15, 41, 66)
    c_blue_primary = RGBColor(2, 132, 199)
    c_blue_dark = RGBColor(30, 58, 138)
    c_blue_light = RGBColor(239, 246, 255)
    c_blue_border = RGBColor(191, 219, 254)
    
    c_red_primary = RGBColor(220, 38, 38)
    c_red_dark = RGBColor(185, 28, 28)
    c_red_light = RGBColor(254, 242, 242)
    c_red_border = RGBColor(254, 202, 202)
    
    c_amber_primary = RGBColor(217, 119, 6)
    c_amber_dark = RGBColor(180, 83, 9)
    c_amber_light = RGBColor(255, 251, 235)
    c_amber_border = RGBColor(253, 230, 138)
    
    c_teal_primary = RGBColor(13, 148, 136)
    c_teal_dark = RGBColor(15, 118, 110)
    c_teal_light = RGBColor(240, 253, 250)
    c_teal_border = RGBColor(153, 246, 228)
    
    c_emerald_primary = RGBColor(5, 150, 105)
    c_emerald_dark = RGBColor(6, 95, 70)
    c_emerald_light = RGBColor(240, 253, 244)
    c_emerald_border = RGBColor(134, 239, 172)
    
    c_purple_primary = RGBColor(124, 58, 237)
    c_purple_dark = RGBColor(109, 40, 217)
    c_purple_light = RGBColor(250, 245, 255)
    c_purple_border = RGBColor(233, 213, 255)
    
    c_slate_dark = RGBColor(30, 41, 59)
    c_slate_muted = RGBColor(100, 116, 139)
    c_slate_light = RGBColor(248, 250, 252)
    c_slate_border = RGBColor(226, 232, 240)
    c_white = RGBColor(255, 255, 255)

    # =========================================================================
    # SLIDE 1: TITLE PAGE
    # =========================================================================
    slide1 = prs.slides[0]
    shapes_to_remove = []
    for s in slide1.shapes:
        if s.name in ["TextBox 10", "Freeform 11", "Freeform 9"]:
            shapes_to_remove.append(s)
        elif s.name == "TextBox 2":
            s.top = Inches(2.05)
            s.height = Inches(1.2)
            
    for s in shapes_to_remove:
        remove_shape(s)
        
    # Project Title & Subtitle Card (Left: 1.0, top: 3.75, width: 9.8, height: 6.4)
    title_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(3.75), Inches(9.8), Inches(6.4))
    title_card.fill.solid()
    title_card.fill.fore_color.rgb = c_slate_light
    title_card.line.color.rgb = c_blue_primary
    title_card.line.width = Pt(1.5)
    
    tf1 = title_card.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.3)
    tf1.margin_right = Inches(0.3)
    tf1.margin_top = Inches(0.25)
    tf1.margin_bottom = Inches(0.2)
    
    # Project Name
    p = tf1.paragraphs[0]
    p.text = "RoadGuard AI"
    p.font.name = "Arial"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = c_navy_dark
    p.alignment = PP_ALIGN.LEFT
    
    # Project Subtitle
    p_sub = tf1.add_paragraph()
    p_sub.text = "ML-Based Road Damage Detection Using PMPML Buses"
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(16)
    p_sub.font.bold = True
    p_sub.font.color.rgb = c_blue_primary
    p_sub.space_after = Pt(12)
    
    meta_items = [
        ("Problem Statement Title", "RoadGuard AI: ML-Based Road Damage Detection and Classification System"),
        ("Theme", "Artificial Intelligence (AI), Smart Mobility and Road Safety"),
        ("Team Name", "404: Ground Not Found"),
        ("College", "PICT (Pune Institute of Computer Technology)"),
        ("Department", "ENTC (Electronics & Telecommunication)")
    ]
    
    for label, val in meta_items:
        p_item = tf1.add_paragraph()
        r1 = p_item.add_run()
        r1.text = f"• {label}: "
        r1.font.name = "Arial"
        r1.font.size = Pt(15)
        r1.font.bold = True
        r1.font.color.rgb = c_slate_dark
        
        r2 = p_item.add_run()
        r2.text = val
        r2.font.name = "Arial"
        r2.font.size = Pt(15)
        r2.font.bold = (label == "Team Name")
        r2.font.color.rgb = c_blue_dark if label == "Team Name" else c_slate_dark
        p_item.space_after = Pt(6)
        
    p_badge = tf1.add_paragraph()
    p_badge.space_before = Pt(8)
    r_b = p_badge.add_run()
    r_b.text = "Track: Hardware Hackathon  |  Platform: Edge-AI & PMPML Transit Fleet"
    r_b.font.name = "Arial"
    r_b.font.size = Pt(11.5)
    r_b.font.bold = True
    r_b.font.color.rgb = c_slate_muted
    
    # Right Side: Realistic PMPML Bus Image
    slide1.shapes.add_picture(pmpml_bus_img, Inches(11.2), Inches(3.75), Inches(7.8), Inches(5.35))
    
    # Caption Box below Bus Image
    caption_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.2), Inches(9.2), Inches(7.8), Inches(0.95))
    caption_box.fill.solid()
    caption_box.fill.fore_color.rgb = c_navy_dark
    caption_box.line.color.rgb = c_blue_primary
    caption_box.line.width = Pt(1.0)
    tf_cap = caption_box.text_frame
    tf_cap.word_wrap = True
    tf_cap.margin_top = Inches(0.12)
    p_cap = tf_cap.paragraphs[0]
    p_cap.text = "PMPML Transit Bus • Real-Time AI Road Vibration Telemetry Scan (Pune)"
    p_cap.font.name = "Arial"
    p_cap.font.size = Pt(12.5)
    p_cap.font.bold = True
    p_cap.font.color.rgb = c_white
    p_cap.alignment = PP_ALIGN.CENTER
    
    p_cap_sub = tf_cap.add_paragraph()
    p_cap_sub.text = "Passive Mobile Road Inspection via Operational Public Transportation Fleet"
    p_cap_sub.font.name = "Arial"
    p_cap_sub.font.size = Pt(10.5)
    p_cap_sub.font.color.rgb = RGBColor(186, 230, 253)
    p_cap_sub.alignment = PP_ALIGN.CENTER

    print("Slide 1 generated.")

    # =========================================================================
    # SLIDE 2: IDEA TITLE / PROPOSED SOLUTION (FLOWCHART & VISUAL COMPARISON)
    # =========================================================================
    slide2 = prs.slides[1]
    for s in list(slide2.shapes):
        if s.name == "TextBox 5":
            remove_shape(s)
            
    # Sub-header: Idea Title & Proposed Solution
    sub_card = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.05), Inches(18.0), Inches(0.85))
    sub_card.fill.solid()
    sub_card.fill.fore_color.rgb = c_blue_light
    sub_card.line.color.rgb = c_blue_primary
    sub_card.line.width = Pt(1.5)
    tf_sub = sub_card.text_frame
    tf_sub.word_wrap = True
    tf_sub.margin_left = Inches(0.2)
    tf_sub.margin_top = Inches(0.06)
    
    p = tf_sub.paragraphs[0]
    r = p.add_run()
    r.text = "Proposed Solution (Describe your Idea/Solution/Prototype)"
    r.font.name = "Arial"
    r.font.size = Pt(19)
    r.font.bold = True
    r.font.underline = True
    r.font.color.rgb = RGBColor(31, 73, 125)
    
    p2 = tf_sub.add_paragraph()
    r2 = p2.add_run()
    r2.text = "Idea Title: "
    r2.font.name = "Arial"
    r2.font.size = Pt(16.5)
    r2.font.bold = True
    r2.font.color.rgb = c_navy_dark
    r3 = p2.add_run()
    r3.text = "RoadGuard AI – Intelligent Road-Vibration Classification System"
    r3.font.name = "Arial"
    r3.font.size = Pt(16.5)
    r3.font.bold = True
    r3.font.color.rgb = c_blue_primary
    
    # 1. Concept Flowchart Banner (top: 3.0, h: 2.0, w: 18.0)
    # Shows Challenge -> Solution -> Innovation Flowchart
    slide2.shapes.add_picture(s2_concept_flowchart_img, Inches(1.0), Inches(3.0), Inches(18.0), Inches(2.0))
    
    # 2. Visual Comparison Cards Infographic (top: 5.15, h: 5.0, w: 18.0)
    # Shows Pothole vs Speed Breaker vs Rough Road geometry & signatures
    slide2.shapes.add_picture(comparison_cards_img, Inches(1.0), Inches(5.15), Inches(18.0), Inches(4.95))
    print("Slide 2 generated.")

    # =========================================================================
    # SLIDE 3: TECHNICAL APPROACH (FLOWCHARTS + PROTOTYPE PHOTO SHOWCASE)
    # =========================================================================
    slide3 = prs.slides[2]
    for s in list(slide3.shapes):
        if s.name == "TextBox 5":
            remove_shape(s)
            
    # 1. Top Section: 7-Step Working Methodology Flowchart banner (top: 2.05, h: 2.15, w: 18.0)
    slide3.shapes.add_picture(flowchart_img, Inches(1.0), Inches(2.05), Inches(18.0), Inches(2.15))
    
    # Bottom Section: Three Columns (y: 4.35 to 10.15)
    # COLUMN 1: DEDICATED EXPECTED PROTOTYPE SHOWCASE CARD (w: 6.0 in)
    proto_container = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.35), Inches(6.0), Inches(5.8))
    proto_container.fill.solid()
    proto_container.fill.fore_color.rgb = c_slate_light
    proto_container.line.color.rgb = c_blue_primary
    proto_container.line.width = Pt(1.5)
    
    # Header bar
    proto_hdr = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.35), Inches(6.0), Inches(0.42))
    proto_hdr.fill.solid()
    proto_hdr.fill.fore_color.rgb = c_navy_dark
    proto_hdr.line.color.rgb = c_navy_dark
    tf_ph = proto_hdr.text_frame
    p_ph = tf_ph.paragraphs[0]
    p_ph.text = "EXPECTED PROTOTYPE SETUP (HARDWARE)"
    p_ph.font.name = "Arial"
    p_ph.font.size = Pt(12)
    p_ph.font.bold = True
    p_ph.font.color.rgb = c_white
    p_ph.alignment = PP_ALIGN.CENTER
    
    # Dedicated Prototype Photo Placeholder / Slot (w: 5.7, h: 3.55)
    slide3.shapes.add_picture(proto_placeholder_img, Inches(1.15), Inches(4.85), Inches(5.7), Inches(3.55))
    
    # Concise Hardware Specs Box below photo (top: 8.50, h: 1.55)
    hw_spec_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.15), Inches(8.50), Inches(5.7), Inches(1.55))
    hw_spec_box.fill.solid()
    hw_spec_box.fill.fore_color.rgb = c_white
    hw_spec_box.line.color.rgb = c_slate_border
    hw_spec_box.line.width = Pt(1.0)
    tf_hws = hw_spec_box.text_frame
    tf_hws.word_wrap = True
    tf_hws.margin_left = Inches(0.12)
    tf_hws.margin_right = Inches(0.12)
    tf_hws.margin_top = Inches(0.06)
    tf_hws.margin_bottom = Inches(0.06)
    
    specs = [
        ("Sensor:", "MPU6050 6-Axis IMU (Acc ±2g to ±16g, 3-Axis Gyro)"),
        ("Microcontroller:", "ESP32 / Arduino Uno (I2C Bus @ 50–100Hz)"),
        ("Enclosure:", "Rigid IP65 Casing with Vibration-Damped Chassis Mount"),
        ("Telemetry & Power:", "5V/3.3V Regulated; Optional NEO-6M GPS & GSM"),
        ("Firmware Stack:", "Arduino C++, Scikit-learn Random Forest Model")
    ]
    for idx, (label, val) in enumerate(specs):
        p_sp = tf_hws.paragraphs[0] if idx == 0 else tf_hws.add_paragraph()
        r1 = p_sp.add_run()
        r1.text = f"• {label} "
        r1.font.name = "Arial"
        r1.font.size = Pt(9.2)
        r1.font.bold = True
        r1.font.color.rgb = c_navy_dark
        r2 = p_sp.add_run()
        r2.text = val
        r2.font.name = "Arial"
        r2.font.size = Pt(8.8)
        r2.font.color.rgb = c_slate_dark
        p_sp.space_after = Pt(1.0)

    # COLUMN 2: VIBRATION SIGNATURE ANALYSIS & WAVEFORMS (w: 7.2 in)
    wave_container = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.15), Inches(4.35), Inches(7.2), Inches(5.8))
    wave_container.fill.solid()
    wave_container.fill.fore_color.rgb = c_slate_light
    wave_container.line.color.rgb = c_blue_primary
    wave_container.line.width = Pt(1.5)
    
    # Header
    wave_hdr = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.15), Inches(4.35), Inches(7.2), Inches(0.42))
    wave_hdr.fill.solid()
    wave_hdr.fill.fore_color.rgb = c_blue_primary
    wave_hdr.line.color.rgb = c_blue_primary
    tf_wh = wave_hdr.text_frame
    p_wh = tf_wh.paragraphs[0]
    p_wh.text = "VIBRATION SIGNATURE ANALYSIS (Z-AXIS ACCEL)"
    p_wh.font.name = "Arial"
    p_wh.font.size = Pt(11.5)
    p_wh.font.bold = True
    p_wh.font.color.rgb = c_white
    p_wh.alignment = PP_ALIGN.CENTER
    
    # Waveform Image (w: 7.0, h: 2.85)
    slide3.shapes.add_picture(waveforms_img, Inches(7.25), Inches(4.85), Inches(7.0), Inches(2.85))
    
    # Analysis Box below waveforms (top: 7.8, h: 2.25)
    tf_wav_box = add_header_card(slide3, 7.25, 7.8, 7.0, 0.35, 2.25,
                                 "ILLUSTRATIVE WAVEFORM DISCRIMINATION PATTERNS",
                                 c_navy_dark, c_white, c_blue_border, title_size=10.5)
    wav_analysis = [
        ("Pothole (Cavity):", "Sharp negative drop followed by high rebound spike (illustrative ~150–250ms transient). Characterized by high kurtosis and high delta-g relative to baseline.", c_red_dark),
        ("Speed Breaker (Hump):", "Smooth, symmetrical low-frequency elevation (illustrative ~600–900ms duration at ~20–30 km/h). Distinct low-frequency sinusoidal profile without sharp shock.", c_amber_primary),
        ("Rough Road (Texture):", "Continuous irregular vibration chatter (illustrative ±0.4g spread around 1.0g gravity baseline). High variance without isolated high-energy peaks.", c_blue_primary),
        ("* Calibration Note:", "Waveforms and thresholds shown are illustrative benchmarks from published research; final operational thresholds will be empirically calibrated using onboard bus sensor data.", c_slate_dark)
    ]
    for idx, (head, body, col) in enumerate(wav_analysis):
        p_h = tf_wav_box.paragraphs[0] if idx == 0 else tf_wav_box.add_paragraph()
        r1 = p_h.add_run()
        r1.text = f"• {head} "
        r1.font.name = "Arial"
        r1.font.size = Pt(8.6) if idx == 3 else Pt(9.2)
        r1.font.bold = True
        r1.font.italic = (idx == 3)
        r1.font.color.rgb = col
        r2 = p_h.add_run()
        r2.text = body
        r2.font.name = "Arial"
        r2.font.size = Pt(8.2) if idx == 3 else Pt(8.8)
        r2.font.italic = (idx == 3)
        r2.font.color.rgb = c_slate_dark
        p_h.space_after = Pt(1.5)

    # COLUMN 3: MACHINE LEARNING PIPELINE FLOWCHART & SCOPE (w: 4.5 in)
    ml_container = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(14.5), Inches(4.35), Inches(4.5), Inches(5.8))
    ml_container.fill.solid()
    ml_container.fill.fore_color.rgb = c_purple_light
    ml_container.line.color.rgb = c_purple_primary
    ml_container.line.width = Pt(1.5)
    
    ml_hdr = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(14.5), Inches(4.35), Inches(4.5), Inches(0.42))
    ml_hdr.fill.solid()
    ml_hdr.fill.fore_color.rgb = c_purple_primary
    ml_hdr.line.color.rgb = c_purple_primary
    tf_mlh = ml_hdr.text_frame
    p_mlh = tf_mlh.paragraphs[0]
    p_mlh.text = "5-STEP ML WORKFLOW"
    p_mlh.font.name = "Arial"
    p_mlh.font.size = Pt(11.5)
    p_mlh.font.bold = True
    p_mlh.font.color.rgb = c_white
    p_mlh.alignment = PP_ALIGN.CENTER
    
    # ML Flowchart Image (w: 4.25, h: 3.7)
    slide3.shapes.add_picture(s3_ml_flowchart_img, Inches(14.62), Inches(4.85), Inches(4.25), Inches(3.7))
    
    # Prototype Scope Note below flowchart (top: 8.65, h: 1.4)
    scope_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(14.62), Inches(8.65), Inches(4.25), Inches(1.4))
    scope_box.fill.solid()
    scope_box.fill.fore_color.rgb = c_white
    scope_box.line.color.rgb = c_purple_border
    scope_box.line.width = Pt(1.0)
    tf_sc = scope_box.text_frame
    tf_sc.word_wrap = True
    tf_sc.margin_left = Inches(0.1)
    tf_sc.margin_right = Inches(0.1)
    tf_sc.margin_top = Inches(0.06)
    
    p_sn = tf_sc.paragraphs[0]
    p_sn.text = "Prototype Configuration Note:"
    p_sn.font.name = "Arial"
    p_sn.font.size = Pt(9.5)
    p_sn.font.bold = True
    p_sn.font.color.rgb = c_navy_dark
    
    p_snd = tf_sc.add_paragraph()
    p_snd.text = "Exact ML algorithm, deployment platform (edge MCU vs companion PC), and optional modules depend on final prototype setup."
    p_snd.font.name = "Arial"
    p_snd.font.size = Pt(8.8)
    p_snd.font.color.rgb = c_slate_muted
    p_snd.space_before = Pt(2)
    
    print("Slide 3 generated.")

    # =========================================================================
    # SLIDE 4: FEASIBILITY AND VIABILITY (SUMMARIZED INFOGRAPHIC)
    # =========================================================================
    slide4 = prs.slides[3]
    for s in list(slide4.shapes):
        if s.name == "TextBox 5":
            remove_shape(s)
            
    # Three-part infographic: Feasibility, Challenges, Mitigation
    # Width: 5.7 each, lefts: 1.0, 7.15, 13.3, top: 2.2, height: 7.2
    # Card 1: Feasibility
    tf_feas = add_header_card(slide4, 1.0, 2.2, 5.7, 0.45, 7.15,
                              "1. FEASIBILITY (Why It Works)", 
                              c_emerald_primary, c_emerald_light, c_emerald_border, title_size=12.5)
    feas_items = [
        ("Sub-$2 Sensor Hardware:", "Uses commodity MPU6050 MEMS sensor (sub-$2), keeping build cost negligible."),
        ("Rapid Microcontroller Setup:", "Easily prototyped with ESP32 / Arduino Uno interfaced with a laptop for data capture."),
        ("Lightweight Machine Learning:", "Random Forest runs with tiny compute footprint—zero costly GPU infrastructure needed."),
        ("Proposed Field Testing:", "Planned experimental trials and calibration across diverse Pune road surfaces and PMPML transit routes using labeled sensor runs."),
        ("Modular Architecture:", "Decoupled data collection, preprocessing, and classification for easy incremental upgrades.")
    ]
    for idx, (t, d) in enumerate(feas_items):
        p_t = tf_feas.paragraphs[0] if idx == 0 else tf_feas.add_paragraph()
        r1 = p_t.add_run()
        r1.text = f"[ ✔ ] {t} "
        r1.font.name = "Arial"
        r1.font.size = Pt(11.5)
        r1.font.bold = True
        r1.font.color.rgb = c_emerald_dark
        r2 = p_t.add_run()
        r2.text = d
        r2.font.name = "Arial"
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = c_slate_dark
        p_t.space_after = Pt(10)

    # Card 2: Potential Challenges
    tf_chal = add_header_card(slide4, 7.15, 2.2, 5.7, 0.45, 7.15,
                              "2. POTENTIAL CHALLENGES", 
                              c_amber_primary, c_amber_light, c_amber_border, title_size=12.5)
    chal_items = [
        ("Bus Chassis & Suspension Variance:", "Different bus models and mounting locations produce shifted baseline vibration readings."),
        ("Vibration Waveform Overlap:", "Steep speed breakers crossed at speed can mimic acceleration peaks of shallow potholes."),
        ("Vehicle Speed & Payload Dynamics:", "Fluctuating passenger weight and driving speeds influence raw vibration amplitudes."),
        ("Representative Dataset Needs:", "Accurate ML categorization requires high-quality, ground-truthed training datasets."),
        ("Transit Environmental Shocks:", "Hardware deployed on city buses faces continuous vibration, dust, and temperature shifts.")
    ]
    for idx, (t, d) in enumerate(chal_items):
        p_t = tf_chal.paragraphs[0] if idx == 0 else tf_chal.add_paragraph()
        r1 = p_t.add_run()
        r1.text = f"[ ! ] {t} "
        r1.font.name = "Arial"
        r1.font.size = Pt(11.5)
        r1.font.bold = True
        r1.font.color.rgb = c_amber_dark
        r2 = p_t.add_run()
        r2.text = d
        r2.font.name = "Arial"
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = c_slate_dark
        p_t.space_after = Pt(10)

    # Card 3: Mitigation Strategies
    tf_mit = add_header_card(slide4, 13.3, 2.2, 5.7, 0.45, 7.15,
                             "3. MITIGATION STRATEGIES", 
                             c_blue_primary, c_blue_light, c_blue_border, title_size=12.5)
    mit_items = [
        ("Standardized Rigid Axle Mount:", "Enclose sensor in rigid, damped IP65 casing bolted to the chassis to ensure clean vibration transfer."),
        ("Multi-Speed Calibration Data:", "Collect labeled calibration data across varying speeds (15–50 km/h) and varied pavement textures."),
        ("Multi-Feature Pattern Decision:", "Classify complete waveforms (duration, RMS, kurtosis, peak-to-peak delta) rather than simple thresholds."),
        ("Rigorous Cross-Validation:", "Evaluate model on separate, unseen test routes with k-fold validation before operational claims."),
        ("Adaptive Baseline Filtering:", "Implement dynamic moving-average filtering to compensate for bus payload changes.")
    ]
    for idx, (t, d) in enumerate(mit_items):
        p_t = tf_mit.paragraphs[0] if idx == 0 else tf_mit.add_paragraph()
        r1 = p_t.add_run()
        r1.text = f"[ ★ ] {t} "
        r1.font.name = "Arial"
        r1.font.size = Pt(11.5)
        r1.font.bold = True
        r1.font.color.rgb = c_blue_dark
        r2 = p_t.add_run()
        r2.text = d
        r2.font.name = "Arial"
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = c_slate_dark
        p_t.space_after = Pt(10)

    # Bottom Summary Banner on Slide 4 (top: 9.45, h: 0.8, w: 18.0)
    sum_feas_card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(9.45), Inches(18.0), Inches(0.8))
    sum_feas_card.fill.solid()
    sum_feas_card.fill.fore_color.rgb = c_navy_dark
    sum_feas_card.line.color.rgb = c_blue_primary
    sum_feas_card.line.width = Pt(1.0)
    tf_sfc = sum_feas_card.text_frame
    p_sfc = tf_sfc.paragraphs[0]
    p_sfc.text = "Feasibility Summary: Sub-$15 Hardware BOM  •  High Municipal Scalability  •  Robust Multi-Feature ML Classification"
    p_sfc.font.name = "Arial"
    p_sfc.font.size = Pt(12.5)
    p_sfc.font.bold = True
    p_sfc.font.color.rgb = c_white
    p_sfc.alignment = PP_ALIGN.CENTER
    
    print("Slide 4 generated.")

    # =========================================================================
    # SLIDE 5: IMPACT AND BENEFITS (FLOWCHART ROADMAP + DASHBOARD)
    # =========================================================================
    slide5 = prs.slides[4]
    for s in list(slide5.shapes):
        if s.name == "TextBox 5":
            remove_shape(s)
            
    has_slide_num = any('5' in (sh.text_frame.text if sh.has_text_frame else '') for sh in slide5.shapes)
    if not has_slide_num:
        num_box = slide5.shapes.add_textbox(Inches(18.0), Inches(10.43), Inches(1.0), Inches(0.5))
        tf_n = num_box.text_frame
        p_n = tf_n.paragraphs[0]
        p_n.text = "5"
        p_n.font.name = "Arial"
        p_n.font.size = Pt(18)
        p_n.font.bold = True
        p_n.font.color.rgb = c_slate_muted

    # Left Column: Impact Summary Cards + Roadmap Flowchart (w: 9.8 in)
    # Social Impact (y: 2.2, h: 1.55)
    tf_soc = add_header_card(slide5, 1.0, 2.2, 9.8, 0.35, 1.55,
                             "SOCIAL IMPACT (ROAD SAFETY & COMMUTER WELFARE)",
                             RGBColor(79, 70, 229), RGBColor(245, 243, 255), RGBColor(221, 214, 254), title_size=11)
    p = tf_soc.paragraphs[0]
    p.text = "• Safer Urban Commutes: "
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(67, 56, 202)
    r = p.add_run()
    r.text = "Continuous passive road auditing enables proactive maintenance for safer roads."
    r.font.size = Pt(10.2)
    r.font.bold = False
    r.font.color.rgb = c_slate_dark
    
    p2 = tf_soc.add_paragraph()
    p2.text = "• Prevents Fatal Accidents: "
    p2.font.name = "Arial"
    p2.font.size = Pt(11)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(67, 56, 202)
    r2 = p2.add_run()
    r2.text = "Early detection of hazardous potholes directly prevents two-wheeler skids & vehicle damage."
    r2.font.size = Pt(10.2)
    r2.font.bold = False
    r2.font.color.rgb = c_slate_dark
    p2.space_before = Pt(3)

    # Economic Impact (y: 3.85, h: 1.55)
    tf_eco = add_header_card(slide5, 1.0, 3.85, 9.8, 0.35, 1.55,
                             "ECONOMIC IMPACT (COST EFFICIENCY & RESOURCE ALLOCATION)",
                             c_emerald_primary, c_emerald_light, c_emerald_border, title_size=11)
    p = tf_eco.paragraphs[0]
    p.text = "• Reduces Survey Expenditure: "
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = c_emerald_dark
    r = p.add_run()
    r.text = "Automates audits, drastically cutting manual survey labor & dedicated vehicle fuel costs."
    r.font.size = Pt(10.2)
    r.font.bold = False
    r.font.color.rgb = c_slate_dark
    
    p2 = tf_eco.add_paragraph()
    p2.text = "• Data-Driven Prioritization: "
    p2.font.name = "Arial"
    p2.font.size = Pt(11)
    p2.font.bold = True
    p2.font.color.rgb = c_emerald_dark
    r2 = p2.add_run()
    r2.text = "Provides objective damage records to help municipal authorities triage repair budgets."
    r2.font.size = Pt(10.2)
    r2.font.bold = False
    r2.font.color.rgb = c_slate_dark
    p2.space_before = Pt(3)

    # Environmental Impact (y: 5.5, h: 1.55)
    tf_env = add_header_card(slide5, 1.0, 5.5, 9.8, 0.35, 1.55,
                             "ENVIRONMENTAL & OPERATIONAL IMPACT (SUSTAINABLE MOBILITY)",
                             c_teal_primary, c_teal_light, c_teal_border, title_size=11)
    p = tf_env.paragraphs[0]
    p.text = "• Zero Dedicated Survey Emissions: "
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = c_teal_dark
    r = p.add_run()
    r.text = "Uses scheduled in-service buses—zero extra survey trips or added carbon emissions."
    r.font.size = Pt(10.2)
    r.font.bold = False
    r.font.color.rgb = c_slate_dark
    
    p2 = tf_env.add_paragraph()
    p2.text = "• Non-Intrusive City Coverage: "
    p2.font.name = "Arial"
    p2.font.size = Pt(11)
    p2.font.bold = True
    p2.font.color.rgb = c_teal_dark
    r2 = p2.add_run()
    r2.text = "Runs silently during regular passenger service without traffic disruptions or road closures."
    r2.font.size = Pt(10.2)
    r2.font.bold = False
    r2.font.color.rgb = c_slate_dark
    p2.space_before = Pt(3)

    # Future Scope Roadmap Flowchart Container (y: 7.15, h: 3.1)
    fut_container = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(7.15), Inches(9.8), Inches(3.1))
    fut_container.fill.solid()
    fut_container.fill.fore_color.rgb = c_slate_light
    fut_container.line.color.rgb = c_navy_dark
    fut_container.line.width = Pt(1.5)
    
    fut_hdr = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(7.15), Inches(9.8), Inches(0.35))
    fut_hdr.fill.solid()
    fut_hdr.fill.fore_color.rgb = c_navy_dark
    fut_hdr.line.color.rgb = c_navy_dark
    tf_fh = fut_hdr.text_frame
    p_fh = tf_fh.paragraphs[0]
    p_fh.text = "FUTURE SCOPE & SYSTEM ROADMAP (PROPOSED EXTENSIONS)"
    p_fh.font.name = "Arial"
    p_fh.font.size = Pt(10.5)
    p_fh.font.bold = True
    p_fh.font.color.rgb = c_white
    p_fh.alignment = PP_ALIGN.CENTER
    
    # Roadmap Flowchart Image (w: 9.6, h: 1.45)
    slide5.shapes.add_picture(s5_roadmap_flowchart_img, Inches(1.1), Inches(7.55), Inches(9.6), Inches(1.45))
    
    # Roadmap description bullets below flowchart
    road_desc_box = slide5.shapes.add_textbox(Inches(1.1), Inches(9.05), Inches(9.6), Inches(1.15))
    tf_rd = road_desc_box.text_frame
    tf_rd.word_wrap = True
    tf_rd.margin_top = Inches(0.02)
    
    p_r1 = tf_rd.paragraphs[0]
    p_r1.text = "• Proposed Stages: [1] GPS Geotagging → [2] GIS Municipal Map → [3] Multi-Bus Fleet Fusion → [4] Severity Estimation → [5] Municipal ERP Ticketing."
    p_r1.font.name = "Arial"
    p_r1.font.size = Pt(9.5)
    p_r1.font.bold = True
    p_r1.font.color.rgb = c_blue_dark
    
    p_r2 = tf_rd.add_paragraph()
    p_r2.text = "• Objective: Enable automated work order generation directly into municipal maintenance workflows."
    p_r2.font.name = "Arial"
    p_r2.font.size = Pt(9)
    p_r2.font.color.rgb = c_slate_dark

    # Right Column: Visual Smart-City Dashboard (left: 11.1, w: 7.9, top: 2.2, h: 4.45)
    slide5.shapes.add_picture(dashboard_img, Inches(11.1), Inches(2.2), Inches(7.9), Inches(4.45))
    
    # Overlay Demo Data Badge on Dashboard Image
    badge_demo = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.25), Inches(2.35), Inches(7.6), Inches(0.38))
    badge_demo.fill.solid()
    badge_demo.fill.fore_color.rgb = RGBColor(15, 23, 42)
    badge_demo.line.color.rgb = c_amber_primary
    badge_demo.line.width = Pt(1.2)
    tf_bd = badge_demo.text_frame
    p_bd = tf_bd.paragraphs[0]
    p_bd.text = "⚠ ILLUSTRATIVE DEMO CONCEPT DATA  •  NOT REAL MUNICIPAL TELEMETRY"
    p_bd.font.name = "Arial"
    p_bd.font.size = Pt(9.5)
    p_bd.font.bold = True
    p_bd.font.color.rgb = RGBColor(253, 230, 138)
    p_bd.alignment = PP_ALIGN.CENTER
    
    # Disclaimer Box below Dashboard Image (y: 6.80, h: 3.45)
    tf_disc = add_header_card(slide5, 11.1, 6.80, 7.9, 0.35, 3.45,
                              "PROPOSED MUNICIPAL GIS CONCEPT (ILLUSTRATIVE DEMO DASHBOARD)",
                              c_navy_dark, c_blue_light, c_blue_border, title_size=10.5)
    p_d_head = tf_disc.paragraphs[0]
    p_d_head.text = "Smart-City Fleet Monitoring Concept (Pune Municipal Corporation):"
    p_d_head.font.name = "Arial"
    p_d_head.font.size = Pt(10.5)
    p_d_head.font.bold = True
    p_d_head.font.color.rgb = c_navy_dark
    
    disc_text = (
        "• Illustrative Demo Data Notice: All statistics, report counts (e.g., 128 reports, 42 active buses), incident alerts, and map pin coordinates shown on this interface are synthetic demonstration data for UI/concept visualization only.\n\n"
        "• Prototype Scope Notice: Real-time GPS mapping, automated cloud telemetry, and centralized municipal dashboard integration represent proposed future extensions and are not part of the initial standalone hardware prototype.\n\n"
        "• Note: No guaranteed financial savings, confirmed government integration, or unverified detection accuracies are claimed."
    )
    p_d_body = tf_disc.add_paragraph()
    p_d_body.text = disc_text
    p_d_body.font.name = "Arial"
    p_d_body.font.size = Pt(9.2)
    p_d_body.font.color.rgb = c_slate_dark
    p_d_body.space_before = Pt(3)

    print("Slide 5 generated.")

    # =========================================================================
    # SLIDE 6: RESEARCH AND REFERENCES (SUMMARIZED CARDS)
    # =========================================================================
    slide6 = prs.slides[5]
    for s in list(slide6.shapes):
        if s.name == "TextBox 5":
            remove_shape(s)
            
    # Four Structured Reference Cards arranged in 2x2 grid (top: 2.2, h: 2.75, w: 8.85 each)
    tf_ref1 = add_header_card(slide6, 1.0, 2.2, 8.85, 0.38, 2.75,
                              "1. HARDWARE SPECIFICATION (MPU-6050)",
                              c_navy_dark, c_slate_light, c_slate_border, title_size=12.5)
    p = tf_ref1.paragraphs[0]
    r = p.add_run()
    r.text = "• Primary Source: "
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = c_navy_dark
    r2 = p.add_run()
    r2.text = "TDK InvenSense MPU-6050 Specification (Doc PS-MPU-6000A-00)"
    r2.font.size = Pt(12.5)
    p.space_after = Pt(8)
    
    p = tf_ref1.add_paragraph()
    r = p.add_run()
    r.text = "• Technical Scope: "
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = c_navy_dark
    r2 = p.add_run()
    r2.text = "Integrated 6-axis MEMS MotionTracking (±2g to ±16g Accel, 3-Axis Gyro, I2C interface)."
    r2.font.size = Pt(12.5)
    p.space_after = Pt(8)
    
    p = tf_ref1.add_paragraph()
    r = p.add_run()
    r.text = "• Project Role: "
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = c_navy_dark
    r2 = p.add_run()
    r2.text = "High-rate road vibration capture and transient chassis impact sampling."
    r2.font.size = Pt(12.5)
    r2.font.color.rgb = c_slate_dark

    # Card 2: Scikit-learn Classifier (left: 10.15, top: 2.2)
    tf_ref2 = add_header_card(slide6, 10.15, 2.2, 8.85, 0.38, 2.75,
                              "2. MACHINE LEARNING FRAMEWORK (RANDOM FOREST)",
                              c_navy_dark, c_slate_light, c_slate_border, title_size=12.5)
    p = tf_ref2.paragraphs[0]
    r = p.add_run()
    r.text = "• Primary Source: "
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = c_navy_dark
    r2 = p.add_run()
    r2.text = "Pedregosa et al., \"Scikit-learn: Machine Learning in Python\" (JMLR 2011)"
    r2.font.size = Pt(12.5)
    p.space_after = Pt(8)
    
    p = tf_ref2.add_paragraph()
    r = p.add_run()
    r.text = "• Technical Scope: "
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = c_navy_dark
    r2 = p.add_run()
    r2.text = "Supervised ensemble classification using RandomForestClassifier algorithm."
    r2.font.size = Pt(12.5)
    p.space_after = Pt(8)
    
    p = tf_ref2.add_paragraph()
    r = p.add_run()
    r.text = "• Project Role: "
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = c_navy_dark
    r2 = p.add_run()
    r2.text = "Multi-class categorization (Pothole vs Speed Breaker vs Rough Road) with noise robustness."
    r2.font.size = Pt(12.5)
    r2.font.color.rgb = c_slate_dark

    # Card 3: Model Evaluation (left: 1.0, top: 5.15)
    tf_ref3 = add_header_card(slide6, 1.0, 5.15, 8.85, 0.38, 2.75,
                              "3. VALIDATION METHODOLOGY (EVALUATION PROTOCOLS)",
                              c_navy_dark, c_slate_light, c_slate_border, title_size=12.5)
    p = tf_ref3.paragraphs[0]
    r = p.add_run()
    r.text = "• Primary Source: "
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = c_navy_dark
    r2 = p.add_run()
    r2.text = "Buitinck et al., \"API Design for Machine Learning Software\" (ECML PKDD)"
    r2.font.size = Pt(12.5)
    p.space_after = Pt(8)
    
    p = tf_ref3.add_paragraph()
    r = p.add_run()
    r.text = "• Technical Scope: "
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = c_navy_dark
    r2 = p.add_run()
    r2.text = "Train/test split protocols (train_test_split), k-fold cross-validation, and multi-class metrics."
    r2.font.size = Pt(12.5)
    p.space_after = Pt(8)
    
    p = tf_ref3.add_paragraph()
    r = p.add_run()
    r.text = "• Project Role: "
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = c_navy_dark
    r2 = p.add_run()
    r2.text = "Ensures model generalization across independent test routes; prevents test data leakage."
    r2.font.size = Pt(12.5)
    r2.font.color.rgb = c_slate_dark

    # Card 4: Published Literature (left: 10.15, top: 5.15)
    tf_ref4 = add_header_card(slide6, 10.15, 5.15, 8.85, 0.38, 2.75,
                              "4. PEER-REVIEWED RESEARCH (TRANSIT ROAD SENSING)",
                              c_navy_dark, c_slate_light, c_slate_border, title_size=12.5)
    papers = [
        ("Mednis et al. (IEEE DCOSS 2011):", "Real-time pothole detection using mobile accelerometers."),
        ("De Zoysa et al. (IEEE ICIAfS 2007):", "Public transport bus-based road condition monitoring."),
        ("Mohan et al. (ACM SenSys 2008):", "Nericell: Mobile sensing for road and traffic classification.")
    ]
    for idx, (auth, desc) in enumerate(papers):
        p_p = tf_ref4.paragraphs[0] if idx == 0 else tf_ref4.add_paragraph()
        r1 = p_p.add_run()
        r1.text = f"• {auth} "
        r1.font.name = "Arial"
        r1.font.size = Pt(12.5)
        r1.font.bold = True
        r1.font.color.rgb = c_navy_dark
        r2 = p_p.add_run()
        r2.text = desc
        r2.font.name = "Arial"
        r2.font.size = Pt(12.5)
        r2.font.color.rgb = c_slate_dark
        p_p.space_after = Pt(8)

    # Bottom Project Summary Callout Box (left: 1.0, top: 8.08, w: 18.0, h: 2.05)
    sum_box = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(8.08), Inches(18.0), Inches(2.05))
    sum_box.fill.solid()
    sum_box.fill.fore_color.rgb = c_blue_light
    sum_box.line.color.rgb = c_blue_primary
    sum_box.line.width = Pt(2.0)
    tf_sum = sum_box.text_frame
    tf_sum.word_wrap = True
    tf_sum.margin_left = Inches(0.4)
    tf_sum.margin_right = Inches(0.4)
    tf_sum.margin_top = Inches(0.18)
    
    p_sum_head = tf_sum.paragraphs[0]
    p_sum_head.text = "PROJECT SUMMARY & CORE INNOVATION"
    p_sum_head.font.name = "Arial"
    p_sum_head.font.size = Pt(13.5)
    p_sum_head.font.bold = True
    p_sum_head.font.color.rgb = c_navy_dark
    p_sum_head.alignment = PP_ALIGN.CENTER
    
    p_quote = tf_sum.add_paragraph()
    p_quote.space_before = Pt(5)
    p_quote.text = "“RoadGuard AI combines low-cost vibration sensing and Machine Learning to distinguish potholes from common road irregularities, supporting smarter road inspection.”"
    p_quote.font.name = "Georgia"
    p_quote.font.size = Pt(17.5)
    p_quote.font.bold = True
    p_quote.font.italic = True
    p_quote.font.color.rgb = c_blue_dark
    p_quote.alignment = PP_ALIGN.CENTER
    
    p_track = tf_sum.add_paragraph()
    p_track.space_before = Pt(5)
    p_track.text = "Hardware Hackathon Track: AI, Smart Mobility and Road Safety  |  Target Deployment: Public Transit Bus Fleets (PMPML, Pune)"
    p_track.font.name = "Arial"
    p_track.font.size = Pt(12)
    p_track.font.bold = True
    p_track.font.color.rgb = c_slate_muted
    p_track.alignment = PP_ALIGN.CENTER

    print("Slide 6 generated.")

    # =========================================================================
    # SLIDE 7: DELETE IMPORTANT INSTRUCTIONS SLIDE
    # =========================================================================
    if len(prs.slides) > 6:
        print(f"Removing Slide 7 ('Important Instructions')...")
        rId = prs.slides._sldIdLst[6].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[6]
        print("Slide 7 removed successfully.")

    print(f"Final slide count: {len(prs.slides)} slides.")
    assert len(prs.slides) == 6, f"Expected exactly 6 slides, got {len(prs.slides)}"
    
    prs.save(output_pptx)
    print(f"Successfully saved refined presentation v2 to: {output_pptx}")

if __name__ == "__main__":
    main()
