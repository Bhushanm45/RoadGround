import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def remove_shape(shape):
    sp = shape._element
    sp.getparent().remove(sp)

def add_header_card(slide, left, top, width, header_height, total_height, 
                    title_text, header_bg, body_bg, border_color, title_size=11.5):
    # Main container
    container = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(total_height))
    container.fill.solid()
    container.fill.fore_color.rgb = body_bg
    container.line.color.rgb = border_color
    container.line.width = Pt(1.5)
    
    # Header bar
    header = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(header_height))
    header.fill.solid()
    header.fill.fore_color.rgb = header_bg
    header.line.color.rgb = header_bg
    header.line.width = Pt(1.0)
    
    tf = header.text_frame
    tf.word_wrap = False
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = "Arial"
    p.font.size = Pt(title_size)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    # Body text frame
    body_box = slide.shapes.add_textbox(Inches(left + 0.08), Inches(top + header_height + 0.04), 
                                        Inches(width - 0.16), Inches(total_height - header_height - 0.08))
    tf_body = body_box.text_frame
    tf_body.word_wrap = True
    tf_body.margin_left = Inches(0.12)
    tf_body.margin_right = Inches(0.12)
    tf_body.margin_top = Inches(0.06)
    tf_body.margin_bottom = Inches(0.06)
    return tf_body

def main():
    template_path = "C:/Users/aryan/.gemini/antigravity/scratch/RoadGuard-AI/template.pptx"
    output_pptx = "C:/Users/aryan/.gemini/antigravity/scratch/RoadGuard-AI/RoadGuard_AI_Hardware_Hackathon_Idea.pptx"
    
    prs = Presentation(template_path)
    print(f"Loaded template with {len(prs.slides)} slides.")
    
    # Assets
    asset_dir = "C:/Users/aryan/.gemini/antigravity/scratch/RoadGuard-AI/assets"
    pmpml_bus_img = os.path.join(asset_dir, "pmpml_bus.jpg")
    sensor_img = os.path.join(asset_dir, "mpu6050_sensor.jpg")
    dashboard_img = os.path.join(asset_dir, "dashboard.jpg")
    comparison_cards_img = os.path.join(asset_dir, "comparison_cards.png")
    flowchart_img = os.path.join(asset_dir, "flowchart.png")
    waveforms_img = os.path.join(asset_dir, "waveforms.png")
    
    # Colors
    c_navy_dark = RGBColor(15, 41, 66)
    c_blue_primary = RGBColor(2, 132, 199)
    c_blue_dark = RGBColor(30, 58, 138)
    c_blue_light = RGBColor(239, 246, 255)
    c_blue_border = RGBColor(191, 219, 254)
    
    c_red_primary = RGBColor(220, 38, 38)
    c_red_dark = RGBColor(185, 28, 28)
    c_red_light = RGBColor(254, 242, 242)
    c_red_border = RGBColor(254, 202, 202)
    
    c_amber_primary = RGBColor(217, 119, 6)
    c_amber_dark = RGBColor(180, 83, 9)
    c_amber_light = RGBColor(255, 251, 235)
    c_amber_border = RGBColor(253, 230, 138)
    
    c_teal_primary = RGBColor(13, 148, 136)
    c_teal_dark = RGBColor(15, 118, 110)
    c_teal_light = RGBColor(240, 253, 250)
    c_teal_border = RGBColor(153, 246, 228)
    
    c_emerald_primary = RGBColor(5, 150, 105)
    c_emerald_dark = RGBColor(6, 95, 70)
    c_emerald_light = RGBColor(240, 253, 244)
    c_emerald_border = RGBColor(134, 239, 172)
    
    c_purple_primary = RGBColor(124, 58, 237)
    c_purple_dark = RGBColor(109, 40, 217)
    c_purple_light = RGBColor(250, 245, 255)
    c_purple_border = RGBColor(233, 213, 255)
    
    c_slate_dark = RGBColor(30, 41, 59)
    c_slate_muted = RGBColor(100, 116, 139)
    c_slate_light = RGBColor(248, 250, 252)
    c_slate_border = RGBColor(226, 232, 240)
    c_white = RGBColor(255, 255, 255)

    # =========================================================================
    # SLIDE 1: TITLE PAGE
    # =========================================================================
    slide1 = prs.slides[0]
    # In slide1, remove placeholder shapes: Freeform 11, Freeform 9, TextBox 10
    shapes_to_remove = []
    for s in slide1.shapes:
        if s.name in ["TextBox 10", "Freeform 11", "Freeform 9"]:
            shapes_to_remove.append(s)
        elif s.name == "TextBox 2":
            # Adjust TextBox 2 so it doesn't overlap
            s.top = Inches(2.05)
            s.height = Inches(1.2)
            
    for s in shapes_to_remove:
        remove_shape(s)
        
    # Project Title & Subtitle Card (Left: 1.0, top: 3.75, width: 9.8, height: 6.4)
    title_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(3.75), Inches(9.8), Inches(6.4))
    title_card.fill.solid()
    title_card.fill.fore_color.rgb = c_slate_light
    title_card.line.color.rgb = c_blue_primary
    title_card.line.width = Pt(1.5)
    
    tf1 = title_card.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.3)
    tf1.margin_right = Inches(0.3)
    tf1.margin_top = Inches(0.25)
    tf1.margin_bottom = Inches(0.2)
    
    # Project Name
    p = tf1.paragraphs[0]
    p.text = "RoadGuard AI"
    p.font.name = "Arial"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = c_navy_dark
    p.alignment = PP_ALIGN.LEFT
    
    # Project Subtitle
    p_sub = tf1.add_paragraph()
    p_sub.text = "ML-Based Road Damage Detection Using PMPML Buses"
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(16)
    p_sub.font.bold = True
    p_sub.font.color.rgb = c_blue_primary
    p_sub.space_after = Pt(12)
    
    # Metadata fields
    meta_items = [
        ("Problem Statement Title", "Intelligent Detection and Classification of Road Damage Using Machine Learning"),
        ("Theme", "AI, Smart Mobility and Road Safety"),
        ("Team Name", "[Insert Team Name]"),
        ("College", "[Insert College Name]"),
        ("Department", "[Insert Department Name]")
    ]
    
    for label, val in meta_items:
        p_item = tf1.add_paragraph()
        r1 = p_item.add_run()
        r1.text = f"• {label}: "
        r1.font.name = "Arial"
        r1.font.size = Pt(15.5)
        r1.font.bold = True
        r1.font.color.rgb = c_slate_dark
        
        r2 = p_item.add_run()
        r2.text = val
        r2.font.name = "Arial"
        r2.font.size = Pt(15.5)
        r2.font.bold = (val.startswith("["))
        r2.font.color.rgb = c_blue_dark if val.startswith("[") else c_slate_dark
        p_item.space_after = Pt(6)
        
    p_badge = tf1.add_paragraph()
    p_badge.space_before = Pt(8)
    r_b = p_badge.add_run()
    r_b.text = "Track: Hardware Hackathon  |  Platform: Edge-AI & PMPML Transit Fleet"
    r_b.font.name = "Arial"
    r_b.font.size = Pt(11.5)
    r_b.font.bold = True
    r_b.font.color.rgb = c_slate_muted
    
    # Right Side: Realistic PMPML Bus Image (top: 3.75, h: 5.35, w: 7.8)
    slide1.shapes.add_picture(pmpml_bus_img, Inches(11.2), Inches(3.75), Inches(7.8), Inches(5.35))
    
    # Caption Box below Bus Image (top: 9.2, h: 0.95, w: 7.8)
    caption_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.2), Inches(9.2), Inches(7.8), Inches(0.95))
    caption_box.fill.solid()
    caption_box.fill.fore_color.rgb = c_navy_dark
    caption_box.line.color.rgb = c_blue_primary
    caption_box.line.width = Pt(1.0)
    tf_cap = caption_box.text_frame
    tf_cap.word_wrap = True
    tf_cap.margin_top = Inches(0.12)
    p_cap = tf_cap.paragraphs[0]
    p_cap.text = "PMPML Transit Bus • Real-Time AI Road Vibration Telemetry Scan (Pune)"
    p_cap.font.name = "Arial"
    p_cap.font.size = Pt(12.5)
    p_cap.font.bold = True
    p_cap.font.color.rgb = c_white
    p_cap.alignment = PP_ALIGN.CENTER
    
    p_cap_sub = tf_cap.add_paragraph()
    p_cap_sub.text = "Passive Mobile Road Inspection via Operational Public Transportation Fleet"
    p_cap_sub.font.name = "Arial"
    p_cap_sub.font.size = Pt(10.5)
    p_cap_sub.font.color.rgb = RGBColor(186, 230, 253)
    p_cap_sub.alignment = PP_ALIGN.CENTER

    print("Slide 1 generated.")

    # =========================================================================
    # SLIDE 2: IDEA TITLE / PROPOSED SOLUTION
    # =========================================================================
    slide2 = prs.slides[1]
    for s in list(slide2.shapes):
        if s.name == "TextBox 5":
            remove_shape(s)
            
    # Sub-header: Idea Title & Proposed Solution
    sub_card = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.05), Inches(18.0), Inches(0.9))
    sub_card.fill.solid()
    sub_card.fill.fore_color.rgb = c_blue_light
    sub_card.line.color.rgb = c_blue_primary
    sub_card.line.width = Pt(1.5)
    tf_sub = sub_card.text_frame
    tf_sub.word_wrap = True
    tf_sub.margin_left = Inches(0.2)
    tf_sub.margin_top = Inches(0.06)
    
    p = tf_sub.paragraphs[0]
    r = p.add_run()
    r.text = "Proposed Solution (Describe your Idea/Solution/Prototype)"
    r.font.name = "Arial"
    r.font.size = Pt(19)
    r.font.bold = True
    r.font.underline = True
    r.font.color.rgb = RGBColor(31, 73, 125)
    
    p2 = tf_sub.add_paragraph()
    r2 = p2.add_run()
    r2.text = "Idea Title: "
    r2.font.name = "Arial"
    r2.font.size = Pt(17)
    r2.font.bold = True
    r2.font.color.rgb = c_navy_dark
    r3 = p2.add_run()
    r3.text = "RoadGuard AI – Intelligent Road-Vibration Classification System"
    r3.font.name = "Arial"
    r3.font.size = Pt(17)
    r3.font.bold = True
    r3.font.color.rgb = c_blue_primary
    
    # 3 Cards: Problem, Solution, Innovation (y: 3.08, h: 3.1, w: 5.7 each)
    tf_prob = add_header_card(slide2, 1.0, 3.08, 5.7, 0.42, 3.1, 
                              "1. PROBLEM DEFINITION", c_red_primary, c_red_light, c_red_border, title_size=12.5)
    p = tf_prob.paragraphs[0]
    p.text = "• Safety Risks & Vehicle Wear:"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = c_red_dark
    p_body = tf_prob.add_paragraph()
    p_body.text = "  Potholes cause critical safety hazards, frequent two-wheeler skids, and costly vehicular suspension damage."
    p_body.font.name = "Arial"
    p_body.font.size = Pt(10.5)
    p_body.font.color.rgb = c_slate_dark
    p_body.space_after = Pt(3)
    
    p = tf_prob.add_paragraph()
    p.text = "• Inefficient Manual Road Inspections:"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = c_red_dark
    p_body = tf_prob.add_paragraph()
    p_body.text = "  Traditional road surveys require extensive manual effort, time, and budget, resulting in delayed repairs."
    p_body.font.name = "Arial"
    p_body.font.size = Pt(10.5)
    p_body.font.color.rgb = c_slate_dark
    p_body.space_after = Pt(3)
    
    p = tf_prob.add_paragraph()
    p.text = "• False Pothole Alarms:"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = c_red_dark
    p_body = tf_prob.add_paragraph()
    p_body.text = "  Speed breakers and ordinary rough roads trigger vibrations that cause conventional systems to generate false alarms."
    p_body.font.name = "Arial"
    p_body.font.size = Pt(10.5)
    p_body.font.color.rgb = c_slate_dark

    # Card 2: Proposed Solution
    tf_sol = add_header_card(slide2, 7.15, 3.08, 5.7, 0.42, 3.1, 
                             "2. PROPOSED SOLUTION", c_blue_primary, c_blue_light, c_blue_border, title_size=12.5)
    p = tf_sol.paragraphs[0]
    p.text = "• MPU6050 Transit Telemetry:"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = c_blue_dark
    p_body = tf_sol.add_paragraph()
    p_body.text = "  Mount an MPU6050 accelerometer & gyroscope on a PMPML bus to collect continuous 6-DOF vibration data."
    p_body.font.name = "Arial"
    p_body.font.size = Pt(10.5)
    p_body.font.color.rgb = c_slate_dark
    p_body.space_after = Pt(3)
    
    p = tf_sol.add_paragraph()
    p.text = "• Feature Extraction & ML Classification:"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = c_blue_dark
    p_body = tf_sol.add_paragraph()
    p_body.text = "  Extract peak acceleration, impact duration, and signal variance; train an ML model to categorize road events."
    p_body.font.name = "Arial"
    p_body.font.size = Pt(10.5)
    p_body.font.color.rgb = c_slate_dark
    p_body.space_after = Pt(3)
    
    p = tf_sol.add_paragraph()
    p.text = "• 3-Way Road Event Discrimination:"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = c_blue_dark
    p_body = tf_sol.add_paragraph()
    p_body.text = "  Distinguishes potholes from speed breakers and ordinary rough roads, logging only genuine damage."
    p_body.font.name = "Arial"
    p_body.font.size = Pt(10.5)
    p_body.font.color.rgb = c_slate_dark

    # Card 3: Innovation & Uniqueness
    tf_inn = add_header_card(slide2, 13.3, 3.08, 5.7, 0.42, 3.1, 
                             "3. INNOVATION & UNIQUENESS", c_teal_primary, c_teal_light, c_teal_border, title_size=12.5)
    p = tf_inn.paragraphs[0]
    p.text = "• Public Buses as Mobile Sensors:"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = c_teal_dark
    p_body = tf_inn.add_paragraph()
    p_body.text = "  Leverages regular PMPML routes as passive road-monitoring platforms without dedicated survey vehicles."
    p_body.font.name = "Arial"
    p_body.font.size = Pt(10.5)
    p_body.font.color.rgb = c_slate_dark
    p_body.space_after = Pt(3)
    
    p = tf_inn.add_paragraph()
    p.text = "• Multi-Class Pattern Analysis:"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = c_teal_dark
    p_body = tf_inn.add_paragraph()
    p_body.text = "  Classifies complete waveform dynamics rather than treating every solitary road bump as a pothole."
    p_body.font.name = "Arial"
    p_body.font.size = Pt(10.5)
    p_body.font.color.rgb = c_slate_dark
    p_body.space_after = Pt(3)
    
    p = tf_inn.add_paragraph()
    p.text = "• Low-Cost, Highly Scalable:"
    p.font.name = "Arial"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = c_teal_dark
    p_body = tf_inn.add_paragraph()
    p_body.text = "  Built with affordable MEMS sensors, enabling city-wide municipal deployment at minimal capital cost."
    p_body.font.name = "Arial"
    p_body.font.size = Pt(10.5)
    p_body.font.color.rgb = c_slate_dark

    # Bottom Area: Visual Comparison of 3 Road Events (y: 6.32, h: 3.75, w: 18.0)
    slide2.shapes.add_picture(comparison_cards_img, Inches(1.0), Inches(6.32), Inches(18.0), Inches(3.75))
    print("Slide 2 generated.")

    # =========================================================================
    # SLIDE 3: TECHNICAL APPROACH
    # =========================================================================
    slide3 = prs.slides[2]
    for s in list(slide3.shapes):
        if s.name == "TextBox 5":
            remove_shape(s)
            
    # Flowchart banner at top (y: 2.05, h: 2.15, w: 18.0)
    slide3.shapes.add_picture(flowchart_img, Inches(1.0), Inches(2.05), Inches(18.0), Inches(2.15))
    
    # Bottom Section: Three Columns (y: 4.35 to 10.15)
    # Column 1: Hardware & Software Stack (w: 5.2)
    # Hardware Card
    tf_hw = add_header_card(slide3, 1.0, 4.35, 5.2, 0.38, 2.8,
                            "HARDWARE ARCHITECTURE", c_navy_dark, c_slate_light, c_slate_border, title_size=11.5)
    p = tf_hw.paragraphs[0]
    p.text = "• MPU6050 IMU: 6-axis accelerometer & gyroscope"
    p.font.name = "Arial"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = c_navy_dark
    
    hw_bullets = [
        "  Measures 3-axis acceleration (g) & angular velocity (deg/s).",
        "• Microcontroller: Arduino Uno or ESP32",
        "  Acquires I2C data at 50–100 Hz (prototype-dependent).",
        "• Power & Enclosure: Regulated 5V/3.3V supply",
        "  Rigid chassis mounting with vibration-damping casing.",
        "• Optional Modules: NEO-6M GPS & GSM/Wi-Fi",
        "  Enables optional geotagging and remote data offload."
    ]
    for b in hw_bullets:
        p_b = tf_hw.add_paragraph()
        p_b.text = b
        p_b.font.name = "Arial"
        p_b.font.size = Pt(9.8)
        p_b.font.bold = b.startswith("•")
        p_b.font.color.rgb = c_navy_dark if b.startswith("•") else c_slate_dark
        p_b.space_after = Pt(2)
        
    # Software Card (y: 7.28, h: 2.85)
    tf_sw = add_header_card(slide3, 1.0, 7.28, 5.2, 0.38, 2.85,
                            "SOFTWARE & ANALYTICS STACK", c_navy_dark, c_slate_light, c_slate_border, title_size=11.5)
    sw_bullets = [
        ("• Arduino IDE / Embedded C++", "Sensor sampling, I2C buffers & serial telemetry."),
        ("• Python (NumPy / SciPy)", "Data calibration, noise filtering & window segmentation."),
        ("• Scikit-learn Framework", "Random Forest classifier training & model evaluation."),
        ("• CSV Labeled Pipelines", "Structured dataset collection for supervised training.")
    ]
    for idx, (title, desc) in enumerate(sw_bullets):
        p_t = tf_sw.paragraphs[0] if idx == 0 else tf_sw.add_paragraph()
        p_t.text = title
        p_t.font.name = "Arial"
        p_t.font.size = Pt(10.2)
        p_t.font.bold = True
        p_t.font.color.rgb = c_blue_dark
        
        p_d = tf_sw.add_paragraph()
        p_d.text = f"  {desc}"
        p_d.font.name = "Arial"
        p_d.font.size = Pt(9.6)
        p_d.font.color.rgb = c_slate_dark
        p_d.space_after = Pt(2.5)

    # Column 2: Waveform Visualizations & Analysis (left: 6.45, w: 7.7)
    slide3.shapes.add_picture(waveforms_img, Inches(6.45), Inches(4.35), Inches(7.7), Inches(2.95))
    
    # Analysis Box below waveforms (y: 7.42, h: 2.7)
    tf_wav_box = add_header_card(slide3, 6.45, 7.42, 7.7, 0.38, 2.7,
                                 "VIBRATION SIGNATURE ANALYSIS (Z-AXIS ACCELEROMETER)",
                                 c_blue_primary, c_slate_light, c_blue_border, title_size=11)
    wav_analysis = [
        ("Pothole (Structural Cavity):", "Sharp downward drop (-0.5g) followed by violent rebound shock (> +3.0g). Transient duration: 150–250ms. High crest factor & kurtosis.", c_red_dark),
        ("Speed Breaker (Engineered Hump):", "Smooth, symmetrical low-frequency sinusoidal hump (+1.95g). Sustained duration: 600–900ms. Dominant low frequency (~1.2–2.0 Hz).", c_amber_primary),
        ("Ordinary Rough Road (Texture Noise):", "Continuous stochastic chatter (±0.4g spread around 1.0g gravity baseline). High variance without solitary extreme peaks.", c_blue_primary)
    ]
    for idx, (head, body, col) in enumerate(wav_analysis):
        p_h = tf_wav_box.paragraphs[0] if idx == 0 else tf_wav_box.add_paragraph()
        r1 = p_h.add_run()
        r1.text = f"• {head} "
        r1.font.name = "Arial"
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = col
        
        r2 = p_h.add_run()
        r2.text = body
        r2.font.name = "Arial"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = c_slate_dark
        p_h.space_after = Pt(2.5)

    # Column 3: ML Pipeline & Implementation Scope (left: 14.38, w: 4.62)
    tf_ml = add_header_card(slide3, 14.38, 4.35, 4.62, 0.38, 3.55,
                            "MACHINE LEARNING PIPELINE", c_purple_primary, c_purple_light, c_purple_border, title_size=11.5)
    ml_steps = [
        ("1. Data Collection", "Log labeled 6-DOF sensor streams on designated road test runs."),
        ("2. Signal Preprocessing", "Remove 1g gravity offset, calibrate zero-drift, and apply sliding window."),
        ("3. Feature Extraction", "Calculate peak acceleration, impact duration, RMS, variance, and kurtosis."),
        ("4. Model Training", "Train Random Forest classifier using cross-validation on labeled dataset."),
        ("5. Event Classification", "Classify incoming vibration windows into Pothole, Speed Breaker, or Rough.")
    ]
    for idx, (st_t, st_d) in enumerate(ml_steps):
        p_s = tf_ml.paragraphs[0] if idx == 0 else tf_ml.add_paragraph()
        r1 = p_s.add_run()
        r1.text = f"• {st_t}: "
        r1.font.name = "Arial"
        r1.font.size = Pt(9.8)
        r1.font.bold = True
        r1.font.color.rgb = c_purple_dark
        r2 = p_s.add_run()
        r2.text = st_d
        r2.font.name = "Arial"
        r2.font.size = Pt(9.3)
        r2.font.color.rgb = c_slate_dark
        p_s.space_after = Pt(2)
        
    # Prototype Scope Note (y: 8.05, h: 2.07)
    tf_note = add_header_card(slide3, 14.38, 8.05, 4.62, 0.38, 2.07,
                              "PROTOTYPE IMPLEMENTATION NOTE", c_slate_dark, c_slate_light, c_slate_border, title_size=10.5)
    p_n = tf_note.paragraphs[0]
    p_n.text = "Hardware & ML Deployment Scope:"
    p_n.font.name = "Arial"
    p_n.font.size = Pt(9.8)
    p_n.font.bold = True
    p_n.font.color.rgb = c_navy_dark
    
    p_n_desc = tf_note.add_paragraph()
    p_n_desc.text = "The exact ML model (e.g. Random Forest vs SVM), deployment platform (microcontroller edge inference vs companion computer), and optional hardware modules (GPS / cellular) depend on the specific hardware prototype configuration implemented."
    p_n_desc.font.name = "Arial"
    p_n_desc.font.size = Pt(9.2)
    p_n_desc.font.color.rgb = c_slate_muted
    
    print("Slide 3 generated.")

    # =========================================================================
    # SLIDE 4: FEASIBILITY AND VIABILITY
    # =========================================================================
    slide4 = prs.slides[3]
    for s in list(slide4.shapes):
        if s.name == "TextBox 5":
            remove_shape(s)
            
    # Three-part infographic: Feasibility, Challenges, Mitigation
    # Width: 5.7 each, lefts: 1.0, 7.15, 13.3, top: 2.2, height: 7.95
    # Card 1: Feasibility
    tf_feas = add_header_card(slide4, 1.0, 2.2, 5.7, 0.45, 7.95,
                              "1. FEASIBILITY ANALYSIS (Why It Works)", 
                              c_emerald_primary, c_emerald_light, c_emerald_border, title_size=12.5)
    feas_items = [
        ("Affordable, Accessible Sensor:", "Uses the widely available MPU6050 6-axis MEMS accelerometer/gyroscope module (sub-$2 component cost), minimizing hardware barriers."),
        ("Rapid Microcontroller Prototyping:", "Can be rapidly prototyped with standard microcontrollers (Arduino Uno or ESP32) interfaced with a computer for initial signal capture and model development."),
        ("Lightweight Machine Learning:", "Classical ML classifiers (such as Random Forest) have low memory and computational footprints, executing efficiently without requiring expensive cloud GPU servers."),
        ("Practical Real-Road Testing:", "Can be readily tested and validated on different road surfaces in Pune by collecting and labeling empirical sensor data during passenger runs."),
        ("Modular System Architecture:", "Decouples sensor data collection, signal filtering, and classification, enabling smooth progressive refinement from prototype to deployment.")
    ]
    for idx, (t, d) in enumerate(feas_items):
        p_t = tf_feas.paragraphs[0] if idx == 0 else tf_feas.add_paragraph()
        r1 = p_t.add_run()
        r1.text = f"[ ✔ ] {t}\n"
        r1.font.name = "Arial"
        r1.font.size = Pt(11.8)
        r1.font.bold = True
        r1.font.color.rgb = c_emerald_dark
        
        r2 = p_t.add_run()
        r2.text = f"      {d}"
        r2.font.name = "Arial"
        r2.font.size = Pt(10.8)
        r2.font.color.rgb = c_slate_dark
        p_t.space_after = Pt(8)

    # Card 2: Potential Challenges
    tf_chal = add_header_card(slide4, 7.15, 2.2, 5.7, 0.45, 7.95,
                              "2. POTENTIAL CHALLENGES & RISKS", 
                              c_amber_primary, c_amber_light, c_amber_border, title_size=12.5)
    chal_items = [
        ("Bus & Mounting Position Variances:", "Different bus models, chassis dampening, and sensor mounting positions may produce shifted baseline vibration readings."),
        ("Vibration Signature Overlap:", "Speed breakers crossed at higher speeds and potholes crossed at lower speeds may sometimes generate overlapping peak acceleration amplitudes."),
        ("Vehicle Dynamic Confounders:", "Vehicle cruising speed, passenger payload changes, tire pressure, and suspension condition directly influence measured vibration dynamics."),
        ("Representative Data Dependency:", "Reliable ML classification requires representative, accurately ground-truthed training data across diverse road types and traffic conditions."),
        ("Environmental & Mechanical Wear:", "Hardware deployed on municipal buses is subjected to continuous transit vibrations, dust, road debris, and thermal fluctuations.")
    ]
    for idx, (t, d) in enumerate(chal_items):
        p_t = tf_chal.paragraphs[0] if idx == 0 else tf_chal.add_paragraph()
        r1 = p_t.add_run()
        r1.text = f"[ ! ] {t}\n"
        r1.font.name = "Arial"
        r1.font.size = Pt(11.8)
        r1.font.bold = True
        r1.font.color.rgb = c_amber_dark
        
        r2 = p_t.add_run()
        r2.text = f"      {d}"
        r2.font.name = "Arial"
        r2.font.size = Pt(10.8)
        r2.font.color.rgb = c_slate_dark
        p_t.space_after = Pt(8)

    # Card 3: Mitigation Strategies
    tf_mit = add_header_card(slide4, 13.3, 2.2, 5.7, 0.45, 7.95,
                             "3. MITIGATION STRATEGIES & SOLUTIONS", 
                             c_blue_primary, c_blue_light, c_blue_border, title_size=12.5)
    mit_items = [
        ("Standardized Secure Mounting:", "Mount the sensor securely in a consistent position (e.g. rigid chassis/axle enclosure) to ensure reproducible vibration transfer."),
        ("Multi-Speed Calibration Runs:", "Collect sensor training samples at varied vehicle speeds (15–50 km/h) and across diverse road surfaces to build a speed-robust dataset."),
        ("Multi-Feature Pattern Recognition:", "Use multiple features (impact duration, signal variance, peak-to-peak delta, kurtosis) and full vibration patterns rather than a solitary amplitude threshold."),
        ("Rigorous Empirical Validation:", "Evaluate the classifier on separate, unseen test data using k-fold cross-validation before making accuracy or operational claims."),
        ("Rugged Protective Enclosure:", "House sensitive electronics in an IP65 vibration-damped enclosure to shield the prototype against road debris, moisture, and shocks.")
    ]
    for idx, (t, d) in enumerate(mit_items):
        p_t = tf_mit.paragraphs[0] if idx == 0 else tf_mit.add_paragraph()
        r1 = p_t.add_run()
        r1.text = f"[ ★ ] {t}\n"
        r1.font.name = "Arial"
        r1.font.size = Pt(11.8)
        r1.font.bold = True
        r1.font.color.rgb = c_blue_dark
        
        r2 = p_t.add_run()
        r2.text = f"      {d}"
        r2.font.name = "Arial"
        r2.font.size = Pt(10.8)
        r2.font.color.rgb = c_slate_dark
        p_t.space_after = Pt(8)

    print("Slide 4 generated.")

    # =========================================================================
    # SLIDE 5: IMPACT AND BENEFITS
    # =========================================================================
    slide5 = prs.slides[4]
    for s in list(slide5.shapes):
        if s.name == "TextBox 5":
            remove_shape(s)
            
    # Add slide number '5' if missing in template
    has_slide_num = any('5' in (sh.text_frame.text if sh.has_text_frame else '') for sh in slide5.shapes)
    if not has_slide_num:
        num_box = slide5.shapes.add_textbox(Inches(18.0), Inches(10.43), Inches(1.0), Inches(0.5))
        tf_n = num_box.text_frame
        p_n = tf_n.paragraphs[0]
        p_n.text = "5"
        p_n.font.name = "Arial"
        p_n.font.size = Pt(18)
        p_n.font.bold = True
        p_n.font.color.rgb = c_slate_muted

    # Left Column: Impact & Benefits + Future Scope (left: 1.0, w: 9.8, h: 7.95)
    # Social Impact (y: 2.2, h: 1.65)
    tf_soc = add_header_card(slide5, 1.0, 2.2, 9.8, 0.38, 1.65,
                             "SOCIAL IMPACT (ROAD SAFETY & COMMUTER WELFARE)",
                             RGBColor(79, 70, 229), RGBColor(245, 243, 255), RGBColor(221, 214, 254), title_size=11)
    p = tf_soc.paragraphs[0]
    p.text = "• Supports Safer, Better-Maintained Urban Roads:"
    p.font.name = "Arial"
    p.font.size = Pt(11.2)
    p.font.bold = True
    p.font.color.rgb = RGBColor(67, 56, 202)
    p_d = tf_soc.add_paragraph()
    p_d.text = "  Enables continuous, passive road health monitoring to support proactive municipal maintenance and safer city transit."
    p_d.font.name = "Arial"
    p_d.font.size = Pt(10.2)
    p_d.font.color.rgb = c_slate_dark
    p_d.space_after = Pt(2)
    
    p = tf_soc.add_paragraph()
    p.text = "• Prevents Critical Hazards & Vehicle Accidents:"
    p.font.name = "Arial"
    p.font.size = Pt(11.2)
    p.font.bold = True
    p.font.color.rgb = RGBColor(67, 56, 202)
    p_d = tf_soc.add_paragraph()
    p_d.text = "  Helps identify hazardous potholes early, directly curbing fatal two-wheeler skids, accidents, and severe vehicle damage."
    p_d.font.name = "Arial"
    p_d.font.size = Pt(10.2)
    p_d.font.color.rgb = c_slate_dark

    # Economic Impact (y: 3.95, h: 1.55)
    tf_eco = add_header_card(slide5, 1.0, 3.95, 9.8, 0.38, 1.55,
                             "ECONOMIC IMPACT (COST EFFICIENCY & RESOURCE ALLOCATION)",
                             c_emerald_primary, c_emerald_light, c_emerald_border, title_size=11)
    p = tf_eco.paragraphs[0]
    p.text = "• May Reduce Manual Inspection Effort & Municipal Expenditure:"
    p.font.name = "Arial"
    p.font.size = Pt(11.2)
    p.font.bold = True
    p.font.color.rgb = c_emerald_dark
    p_d = tf_eco.add_paragraph()
    p_d.text = "  Automates road surface audits, decreasing the need for frequent manual survey teams and recurring audit expenses."
    p_d.font.name = "Arial"
    p_d.font.size = Pt(10.2)
    p_d.font.color.rgb = c_slate_dark
    p_d.space_after = Pt(2)
    
    p = tf_eco.add_paragraph()
    p.text = "• Optimizes Repair Prioritization:"
    p.font.name = "Arial"
    p.font.size = Pt(11.2)
    p.font.bold = True
    p.font.color.rgb = c_emerald_dark
    p_d = tf_eco.add_paragraph()
    p_d.text = "  Provides municipal authorities with objective damage records, enabling data-driven prioritization of repair budgets."
    p_d.font.name = "Arial"
    p_d.font.size = Pt(10.2)
    p_d.font.color.rgb = c_slate_dark

    # Environmental & Operational Impact (y: 5.6, h: 1.6)
    tf_env = add_header_card(slide5, 1.0, 5.6, 9.8, 0.38, 1.6,
                             "ENVIRONMENTAL & OPERATIONAL IMPACT (SUSTAINABLE MOBILITY)",
                             c_teal_primary, c_teal_light, c_teal_border, title_size=11)
    p = tf_env.paragraphs[0]
    p.text = "• Leverages Existing Transit Fleets (Zero Added Survey Emissions):"
    p.font.name = "Arial"
    p.font.size = Pt(11.2)
    p.font.bold = True
    p.font.color.rgb = c_teal_dark
    p_d = tf_env.add_paragraph()
    p_d.text = "  Makes use of buses already travelling throughout the city, eliminating the need for dedicated survey vehicle trips and extra fuel."
    p_d.font.name = "Arial"
    p_d.font.size = Pt(10.2)
    p_d.font.color.rgb = c_slate_dark
    p_d.space_after = Pt(2)
    
    p = tf_env.add_paragraph()
    p.text = "• Continuous, Non-Intrusive City-Wide Coverage:"
    p.font.name = "Arial"
    p.font.size = Pt(11.2)
    p.font.bold = True
    p.font.color.rgb = c_teal_dark
    p_d = tf_env.add_paragraph()
    p_d.text = "  Operates silently during regular passenger service without impeding traffic flow or requiring road closures."
    p_d.font.name = "Arial"
    p_d.font.size = Pt(10.2)
    p_d.font.color.rgb = c_slate_dark

    # Future Scope (y: 7.3, h: 2.85)
    tf_fut = add_header_card(slide5, 1.0, 7.3, 9.8, 0.38, 2.85,
                             "FUTURE SCOPE & PROPOSED SYSTEM EXTENSIONS",
                             c_navy_dark, c_slate_light, c_slate_border, title_size=11)
    fut_items = [
        ("GPS-Based Pothole Mapping [Proposed]:", "Integrate GNSS receiver to geotag detected road damages with spatial coordinates."),
        ("Municipal Web Dashboard [Proposed]:", "Develop an interactive GIS dashboard for Pune road authorities to visualize city-wide road health."),
        ("Multi-Bus Fleet Data Aggregation [Proposed]:", "Aggregate sensor streams across multiple buses to verify potholes through spatial consensus voting."),
        ("Road-Damage Severity Estimation [Proposed]:", "Analyze impact energy, duration, and vertical drop depth to estimate structural severity."),
        ("Municipal Workflow Integration [Proposed]:", "Connect detection alerts with municipal maintenance ticketing systems for automated work orders.")
    ]
    for idx, (t, d) in enumerate(fut_items):
        p_t = tf_fut.paragraphs[0] if idx == 0 else tf_fut.add_paragraph()
        r1 = p_t.add_run()
        r1.text = f"• {t} "
        r1.font.name = "Arial"
        r1.font.size = Pt(10.2)
        r1.font.bold = True
        r1.font.color.rgb = c_blue_dark
        r2 = p_t.add_run()
        r2.text = d
        r2.font.name = "Arial"
        r2.font.size = Pt(9.8)
        r2.font.color.rgb = c_slate_dark
        p_t.space_after = Pt(2)

    # Right Column: Visual Smart-City Dashboard (left: 11.1, w: 7.9, top: 2.2, h: 7.95)
    slide5.shapes.add_picture(dashboard_img, Inches(11.1), Inches(2.2), Inches(7.9), Inches(4.85))
    
    # Disclaimer Box below Dashboard Image (y: 7.15, h: 3.0)
    tf_disc = add_header_card(slide5, 11.1, 7.15, 7.9, 0.38, 3.0,
                              "PROPOSED FUTURE DASHBOARD & MUNICIPAL GIS CONCEPT",
                              c_navy_dark, c_blue_light, c_blue_border, title_size=10.5)
    p_d_head = tf_disc.paragraphs[0]
    p_d_head.text = "Smart-City Fleet Monitoring Concept (Pune Municipal Corporation):"
    p_d_head.font.name = "Arial"
    p_d_head.font.size = Pt(10.5)
    p_d_head.font.bold = True
    p_d_head.font.color.rgb = c_navy_dark
    
    disc_text = (
        "• Visual Illustration: The above interface illustrates our proposed future web dashboard for municipal road engineers, featuring live route tracking, pothole heatmaps, and severity distribution.\n\n"
        "• Prototype Scope Notice: Real-time GPS mapping, automated cloud telemetry, and centralized municipal dashboard integration represent proposed future extensions and are not part of the initial standalone hardware prototype.\n\n"
        "• Note: No guaranteed financial savings, confirmed government integration, or unverified detection accuracies are claimed."
    )
    p_d_body = tf_disc.add_paragraph()
    p_d_body.text = disc_text
    p_d_body.font.name = "Arial"
    p_d_body.font.size = Pt(9.2)
    p_d_body.font.color.rgb = c_slate_dark
    p_d_body.space_before = Pt(3)

    print("Slide 5 generated.")

    # =========================================================================
    # SLIDE 6: RESEARCH AND REFERENCES
    # =========================================================================
    slide6 = prs.slides[5]
    for s in list(slide6.shapes):
        if s.name == "TextBox 5":
            remove_shape(s)
            
    # Four Structured Reference Cards arranged in 2x2 grid (top: 2.2, h: 2.75, w: 8.85 each)
    tf_ref1 = add_header_card(slide6, 1.0, 2.2, 8.85, 0.4, 2.75,
                              "1. HARDWARE SPECIFICATION: MPU-6050 MOTIONTRACKING",
                              c_navy_dark, c_slate_light, c_slate_border, title_size=11.5)
    p = tf_ref1.paragraphs[0]
    r = p.add_run()
    r.text = "• Citation: "
    r.font.bold = True
    r.font.size = Pt(11)
    r2 = p.add_run()
    r2.text = "TDK InvenSense, \"MPU-6000 and MPU-6050 Product Specification Revision 3.4\", Document No. PS-MPU-6000A-00, 2013."
    r2.font.size = Pt(10.5)
    
    p = tf_ref1.add_paragraph()
    r = p.add_run()
    r.text = "• Technical Focus: "
    r.font.bold = True
    r.font.size = Pt(11)
    r2 = p.add_run()
    r2.text = "Integrated 6-axis MEMS MotionTracking device with embedded 3-axis accelerometer (±2g to ±16g range), 3-axis gyroscope, digital-output I2C bus interface, and low-power processing."
    r2.font.size = Pt(10.5)
    
    p = tf_ref1.add_paragraph()
    r = p.add_run()
    r.text = "• Verifiable URL: "
    r.font.bold = True
    r.font.size = Pt(11)
    r2 = p.add_run()
    r2.text = "https://invensense.tdk.com/products/motion-tracking/6-axis/mpu-6050/"
    r2.font.size = Pt(10)
    r2.font.color.rgb = c_blue_primary

    # Card 2: Scikit-learn Classifier (left: 10.15, top: 2.2)
    tf_ref2 = add_header_card(slide6, 10.15, 2.2, 8.85, 0.4, 2.75,
                              "2. MACHINE LEARNING FRAMEWORK: SCIKIT-LEARN RANDOM FOREST",
                              c_navy_dark, c_slate_light, c_slate_border, title_size=11.5)
    p = tf_ref2.paragraphs[0]
    r = p.add_run()
    r.text = "• Citation: "
    r.font.bold = True
    r.font.size = Pt(11)
    r2 = p.add_run()
    r2.text = "Pedregosa, F., Varoquaux, G., Gramfort, A., et al. (2011). \"Scikit-learn: Machine Learning in Python.\" Journal of Machine Learning Research (JMLR), Vol. 12, pp. 2825–2830."
    r2.font.size = Pt(10.5)
    
    p = tf_ref2.add_paragraph()
    r = p.add_run()
    r.text = "• Technical Focus: "
    r.font.bold = True
    r.font.size = Pt(11)
    r2 = p.add_run()
    r2.text = "Ensemble learning via RandomForestClassifier; bagging decision trees for multi-class vibration feature classification, feature importance ranking, and noise robustness."
    r2.font.size = Pt(10.5)
    
    p = tf_ref2.add_paragraph()
    r = p.add_run()
    r.text = "• Verifiable URL: "
    r.font.bold = True
    r.font.size = Pt(11)
    r2 = p.add_run()
    r2.text = "https://scikit-learn.org/stable/modules/ensemble.html#forests-of-randomized-trees"
    r2.font.size = Pt(10)
    r2.font.color.rgb = c_blue_primary

    # Card 3: Model Evaluation (left: 1.0, top: 5.15)
    tf_ref3 = add_header_card(slide6, 1.0, 5.15, 8.85, 0.4, 2.75,
                              "3. METHODOLOGY: MODEL EVALUATION & VALIDATION PROTOCOLS",
                              c_navy_dark, c_slate_light, c_slate_border, title_size=11.5)
    p = tf_ref3.paragraphs[0]
    r = p.add_run()
    r.text = "• Citation: "
    r.font.bold = True
    r.font.size = Pt(11)
    r2 = p.add_run()
    r2.text = "Buitinck, L., Louppe, G., Blondel, M., et al. (2013). \"API design for machine learning software: experiences from the scikit-learn project.\" ECML PKDD Workshop."
    r2.font.size = Pt(10.5)
    
    p = tf_ref3.add_paragraph()
    r = p.add_run()
    r.text = "• Technical Focus: "
    r.font.bold = True
    r.font.size = Pt(11)
    r2 = p.add_run()
    r2.text = "Train/test split protocols (train_test_split), k-fold cross-validation, hyperparameter tuning, and multi-class performance evaluation metrics (Precision, Recall, F1-Score, Confusion Matrix)."
    r2.font.size = Pt(10.5)
    
    p = tf_ref3.add_paragraph()
    r = p.add_run()
    r.text = "• Verifiable URL: "
    r.font.bold = True
    r.font.size = Pt(11)
    r2 = p.add_run()
    r2.text = "https://scikit-learn.org/stable/modules/cross_validation.html"
    r2.font.size = Pt(10)
    r2.font.color.rgb = c_blue_primary

    # Card 4: Published Literature (left: 10.15, top: 5.15)
    tf_ref4 = add_header_card(slide6, 10.15, 5.15, 8.85, 0.4, 2.75,
                              "4. PEER-REVIEWED RESEARCH: ACCELEROMETER ROAD SENSING",
                              c_navy_dark, c_slate_light, c_slate_border, title_size=11.5)
    p = tf_ref4.paragraphs[0]
    r = p.add_run()
    r.text = "• Key Research Literature: "
    r.font.bold = True
    r.font.size = Pt(11)
    
    papers = [
        "Mednis, A., Strazdins, G., Zviedris, R., Kanonirs, G., & Selavo, L. (2011). \"Real time pothole detection using Android smartphones with accelerometers.\" IEEE DCOSS. DOI: 10.1109/DCOSS.2011.5982206",
        "De Zoysa, K., Kasthurirathna, D., & Keppitiyagama, C. (2007). \"A Public Transport System Based Pothole Detection and Road Condition Monitoring System.\" IEEE ICIAfS.",
        "Mohan, P., Padmanabhan, V. N., & Ramjee, R. (2008). \"Nericell: rich monitoring of road and traffic conditions using mobile smartphones.\" ACM SenSys '08. DOI: 10.1145/1460412.1460444"
    ]
    for pap in papers:
        p_p = tf_ref4.add_paragraph()
        p_p.text = f"• {pap}"
        p_p.font.name = "Arial"
        p_p.font.size = Pt(9.6)
        p_p.font.color.rgb = c_slate_dark
        p_p.space_after = Pt(2)

    # Bottom Project Summary Callout Box (left: 1.0, top: 8.08, w: 18.0, h: 2.05)
    sum_box = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(8.08), Inches(18.0), Inches(2.05))
    sum_box.fill.solid()
    sum_box.fill.fore_color.rgb = c_blue_light
    sum_box.line.color.rgb = c_blue_primary
    sum_box.line.width = Pt(2.0)
    tf_sum = sum_box.text_frame
    tf_sum.word_wrap = True
    tf_sum.margin_left = Inches(0.4)
    tf_sum.margin_right = Inches(0.4)
    tf_sum.margin_top = Inches(0.18)
    
    p_sum_head = tf_sum.paragraphs[0]
    p_sum_head.text = "PROJECT SUMMARY & CORE INNOVATION"
    p_sum_head.font.name = "Arial"
    p_sum_head.font.size = Pt(13.5)
    p_sum_head.font.bold = True
    p_sum_head.font.color.rgb = c_navy_dark
    p_sum_head.alignment = PP_ALIGN.CENTER
    
    p_quote = tf_sum.add_paragraph()
    p_quote.space_before = Pt(5)
    p_quote.text = "“RoadGuard AI combines low-cost vibration sensing and Machine Learning to distinguish potholes from common road irregularities, supporting smarter road inspection.”"
    p_quote.font.name = "Georgia"
    p_quote.font.size = Pt(17.5)
    p_quote.font.bold = True
    p_quote.font.italic = True
    p_quote.font.color.rgb = c_blue_dark
    p_quote.alignment = PP_ALIGN.CENTER
    
    p_track = tf_sum.add_paragraph()
    p_track.space_before = Pt(5)
    p_track.text = "Hardware Hackathon Track: AI, Smart Mobility and Road Safety  |  Target Deployment: Public Transit Bus Fleets (PMPML, Pune)"
    p_track.font.name = "Arial"
    p_track.font.size = Pt(12)
    p_track.font.bold = True
    p_track.font.color.rgb = c_slate_muted
    p_track.alignment = PP_ALIGN.CENTER

    print("Slide 6 generated.")

    # =========================================================================
    # SLIDE 7: DELETE IMPORTANT INSTRUCTIONS SLIDE
    # =========================================================================
    if len(prs.slides) > 6:
        print(f"Removing Slide 7 ('Important Instructions')...")
        rId = prs.slides._sldIdLst[6].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[6]
        print("Slide 7 removed successfully.")

    print(f"Final slide count: {len(prs.slides)} slides.")
    assert len(prs.slides) == 6, f"Expected exactly 6 slides, got {len(prs.slides)}"
    
    prs.save(output_pptx)
    print(f"Successfully saved refined presentation to: {output_pptx}")

if __name__ == "__main__":
    main()
